"""Export Excel d'un résultat de requête (brief §8) : synthèse + données détaillées + définitions."""

import io
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .service import Resultat
from .statut import statut

MARINE = "1F2A44"
ENTETE = Font(bold=True, color="FFFFFF")
FOND_ENTETE = PatternFill("solid", fgColor=MARINE)
FONDS_STATUT = {
    "alerte": PatternFill("solid", fgColor="FFE8B3"),
    "critique": PatternFill("solid", fgColor="F8C4C4"),
}
FORMAT_PERIODE = {"heure": "dd/mm/yyyy hh\\h", "jour": "dd/mm/yyyy", "semaine": "dd/mm/yyyy", "mois": "mm/yyyy"}


def _valeur(v):
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return None
    return float(v)


def _ajouter(ws, valeurs):
    """Ajoute une ligne ; un texte commençant par « = » reste du texte (pas de formule)."""
    ws.append(valeurs)
    for c in ws[ws.max_row]:
        if isinstance(c.value, str) and c.value.startswith("="):
            c.data_type = "s"


def _entete(ws, titres):
    _ajouter(ws, titres)
    for c in ws[ws.max_row]:
        c.font, c.fill = ENTETE, FOND_ENTETE
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = f"A{ws.max_row + 1}"


def _largeurs(ws, largeurs):
    for i, l in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = l


def _cellules_kpi(ws, kpis, valeurs, premiere_col):
    ligne = ws.max_row
    for j, (k, v) in enumerate(zip(kpis, valeurs)):
        c = ws.cell(row=ligne, column=premiere_col + j)
        c.number_format = "0.00"
        fond = FONDS_STATUT.get(statut(k, v))
        if fond:
            c.fill = fond


def _titres_kpi(kpis):
    return [f"{k.libelle} ({k.unite})" + (" ≈" if k.qualite == "approx" else "") for k in kpis]


def parametres(resultat: Resultat) -> list[tuple[str, str]]:
    """Description de la requête (onglet Paramètres, page de titre du PowerPoint)."""
    req = resultat.requete
    return [
        ("Technologies", ", ".join(req.techno)),
        ("Périmètre", f"{req.perimetre.type} : {', '.join(req.perimetre.valeurs) or 'tout le réseau'}"),
        ("Période", f"du {req.periode.debut:%d/%m/%Y} au {req.periode.fin:%d/%m/%Y}"),
        ("Fenêtre horaire", req.fenetre_horaire),
        ("Granularité temporelle", req.granularite_temps),
        ("Granularité spatiale", req.granularite_espace),
    ]


def construire(resultat: Resultat) -> bytes:
    req = resultat.requete
    wb = Workbook()

    ws = wb.active
    ws.title = "Paramètres"
    lignes = [*parametres(resultat), ("KPI", ", ".join(req.kpis))]
    for r in resultat.par_techno:
        lignes.append((f"Cellules {r.techno}", r.cellules_demandees))
        if r.cellules_sans_donnees:
            lignes.append((f"Cellules {r.techno} sans donnée", ", ".join(r.cellules_sans_donnees)))
    for a in resultat.avertissements:
        lignes.append(("Avertissement", a))
    for l in lignes:
        _ajouter(ws, l)
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    _largeurs(ws, [28, 90])

    for r in resultat.par_techno:
        titres = _titres_kpi(r.kpis)

        ws = wb.create_sheet(f"Synthèse {r.techno}")
        _entete(ws, ["Entité", *titres])
        vals = [_valeur(r.synthese_globale.get(k.code)) for k in r.kpis]
        _ajouter(ws, ["Ensemble du périmètre", *vals])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        _cellules_kpi(ws, r.kpis, vals, 2)
        for entite, valeurs in r.synthese.iterrows():
            vals = [_valeur(valeurs[k.code]) for k in r.kpis]
            _ajouter(ws, [entite, *vals])
            _cellules_kpi(ws, r.kpis, vals, 2)
        _largeurs(ws, [30, *[18] * len(titres)])

        ws = wb.create_sheet(f"Données {r.techno}")
        _entete(ws, ["Période", "Entité", *titres])
        for (periode, entite), valeurs in r.table.iterrows():
            vals = [_valeur(valeurs[k.code]) for k in r.kpis]
            _ajouter(ws, [periode.to_pydatetime(), entite, *vals])
            ws.cell(row=ws.max_row, column=1).number_format = FORMAT_PERIODE[req.granularite_temps]
            _cellules_kpi(ws, r.kpis, vals, 3)
        _largeurs(ws, [18, 30, *[18] * len(titres)])
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet("Définitions")
    _entete(ws, ["Code", "Libellé", "Unité", "Numérateur", "Dénominateur", "Facteur",
                 "Seuil alerte", "Seuil critique", "Qualité", "Note"])
    for r in resultat.par_techno:
        for k in r.kpis:
            _ajouter(ws, [k.code, k.libelle, k.unite, k.numerateur, k.denominateur or "", k.facteur,
                       k.seuils.alerte, k.seuils.critique, k.qualite, k.note or ""])
    _largeurs(ws, [22, 30, 10, 40, 40, 9, 12, 12, 12, 50])

    flux = io.BytesIO()
    wb.save(flux)
    return flux.getvalue()
