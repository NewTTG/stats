"""Génération PowerPoint (python-pptx), brief §8.

Avec un modèle (``settings.PPTX_MODELE``, ex. charte Helia PRO) : ses dispositions
(``settings.PPTX_DISPOSITIONS``) et les couleurs de son thème sont utilisées, ses
diapositives d'exemple sont ignorées et les dispositions non utilisées retirées du
fichier produit (elles portent l'essentiel des images du modèle).
Sans modèle : charte bleu marine / ambre dessinée par le code.
"""

import colorsys
import math
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.utils import timezone
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Emu, Inches, Pt

BLANC = RGBColor(0xFF, 0xFF, 0xFF)
FOND_STATUT = {"alerte": RGBColor(0xFF, 0xF1, 0xCC), "critique": RGBColor(0xF8, 0xD0, 0xCC)}
LIGNES_PAR_TABLEAU = 14
DISPOSITIONS = {
    "titre": ["Titre", "Diapositive de titre", "Title Slide"],
    "contenu": ["Titre seul", "Title Only"],
    "fin": [],
}
TYPES_TITRE = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.SUBTITLE}
TYPES_PIED = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}


@dataclass
class Charte:
    principale: str  # série « événement », en-têtes de tableau, titres dessinés
    secondaire: str  # série « référence »
    gris: str  # bande min / max de la référence
    texte: str
    series: list[str]  # une couleur par entité (rapport de requête)


CHARTE_DEFAUT = Charte(principale="14284B", secondaire="F0A500", gris="8A94A6", texte="14284B",
                       series=["14284B", "F0A500", "2E86AB", "C0392B", "27AE60", "8E44AD", "7F8C8D", "D35400"])


def _hls(hexa: str):
    r, g, b = (int(hexa[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return colorsys.rgb_to_hls(r, g, b)


def charte_du_theme(xml: bytes) -> Charte:
    """Couleurs d'un thème Office : la couleur d'accent la plus saturée devient la
    principale, les gris (peu saturés) servent pour la référence et la bande."""
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    racine = etree.fromstring(xml)
    accents = [v for i in range(1, 7)
               for v in racine.xpath(f"//a:clrScheme/a:accent{i}/a:srgbClr/@val", namespaces=ns)]
    if len(accents) < 2:
        return CHARTE_DEFAUT
    saturees = sorted((a for a in accents if _hls(a)[2] >= 0.35), key=lambda a: -_hls(a)[2])
    gris = sorted((a for a in accents if _hls(a)[2] < 0.35 and _hls(a)[1] < 0.85), key=lambda a: _hls(a)[1])
    principale = saturees[0] if saturees else accents[0]
    secondaire = gris[0] if gris else (saturees[1] if len(saturees) > 1 else accents[1])
    milieu = gris[len(gris) // 2] if len(gris) > 1 else "A3A9AE"
    autres = [a for a in saturees + gris if a not in (principale, secondaire)]
    return Charte(principale=principale, secondaire=secondaire, gris=milieu, texte=secondaire,
                  series=[principale, secondaire, *autres, *CHARTE_DEFAUT.series[2:]])


def nombre(v, decimales=2) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return "—"
    return f"{v:,.{decimales}f}".replace(",", " ").replace(".", ",")


def pourcentage(v) -> str:
    return "—" if v is None else f"{v:+.1f} %".replace(".", ",").replace("-", "−")


def _rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa)


class Deck:
    def __init__(self):
        modele = getattr(settings, "PPTX_MODELE", None)
        self.avec_modele = bool(modele) and Path(modele).exists()
        noms = {**DISPOSITIONS, **getattr(settings, "PPTX_DISPOSITIONS", {})}
        if self.avec_modele:
            self.prs = Presentation(modele)
            for sld in list(self.prs.slides._sldIdLst):  # diapositives d'exemple du modèle
                self.prs.part.drop_rel(sld.rId)
                self.prs.slides._sldIdLst.remove(sld)
            self.dispositions = {role: self._chercher(n) for role, n in noms.items()}
            self._elaguer()
            theme = self.prs.slide_master.part.part_related_by(RT.THEME).blob
            self.charte = charte_du_theme(theme)
        else:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
            self.dispositions = dict.fromkeys(noms)
            self.charte = CHARTE_DEFAUT
        self.largeur, self.hauteur = self.prs.slide_width, self.prs.slide_height
        self.marge = Inches(0.5)
        self.haut_contenu, self.bas_contenu = Inches(1.35), self.hauteur - Inches(0.4)
        if self.dispositions["contenu"] is not None:
            self._zone_contenu(self.dispositions["contenu"])

    # --------------------------------------------------------------- modèle
    def _chercher(self, noms: list[str]):
        par_nom = {l.name.strip().lower(): l for l in self.prs.slide_layouts}
        return next((par_nom[n.lower()] for n in noms if n.lower() in par_nom), None)

    def _elaguer(self):
        """Retire les dispositions inutilisées : le fichier ne garde que leurs images."""
        gardees = {id(l) for l in self.dispositions.values() if l is not None}
        for master in self.prs.slide_masters:
            for layout in list(master.slide_layouts):
                if id(layout) not in gardees and len(master.slide_layouts) > 1:
                    master.slide_layouts.remove(layout)

    @staticmethod
    def _titre_de(conteneur):
        """Espace réservé portant le titre : le titre déclaré, sinon le texte le plus haut."""
        candidats = [p for p in conteneur.placeholders if p.placeholder_format.type in TYPES_TITRE]
        titres = [p for p in candidats if p.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)]
        return (titres or sorted(candidats, key=lambda p: (p.top, p.left)) or [None])[0]

    def _zone_contenu(self, layout):
        titre = self._titre_de(layout)
        if titre is not None:
            self.haut_contenu = titre.top + titre.height + Inches(0.25)
        # Bas de la zone utile : pied de page, ou bandeau décoratif pleine largeur en bas.
        limites = [p.top for p in layout.placeholders if p.placeholder_format.type in TYPES_PIED]
        limites += [sh.top for sh in layout.shapes if not sh.is_placeholder and sh.top > self.hauteur / 2
                    and sh.width >= 0.8 * self.largeur]
        if limites:
            self.bas_contenu = min(limites) - Inches(0.15)

    @staticmethod
    def _vider(slide, garder=()):
        for ph in list(slide.placeholders):
            if ph not in garder:
                ph._element.getparent().remove(ph._element)

    # --------------------------------------------------------------- dispositions
    def _diapo(self, titre: str):
        layout = self.dispositions["contenu"]
        if layout is not None:
            slide = self.prs.slides.add_slide(layout)
            ph = self._titre_de(slide)
            if ph is not None:
                ph.text = titre
                if len(titre) > 30:  # le titre doit tenir sur une ligne, au-dessus du contenu
                    ph.text_frame.paragraphs[0].font.size = Pt(24 if len(titre) <= 42 else 20)
                self._vider(slide, garder=[ph])
                return slide
            self._vider(slide)
        else:
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[6 if len(self.prs.slide_layouts) > 6 else -1])
        principale = _rgb(self.charte.principale)
        bandeau = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.largeur, Inches(1.0))
        bandeau.fill.solid()
        bandeau.fill.fore_color.rgb = principale
        bandeau.line.fill.background()
        filet = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.0), self.largeur, Inches(0.06))
        filet.fill.solid()
        filet.fill.fore_color.rgb = _rgb(self.charte.secondaire)
        filet.line.fill.background()
        self._texte(slide, titre, self.marge, Inches(0.2), self.largeur - 2 * self.marge, Inches(0.65),
                    taille=26, couleur=BLANC, gras=True)
        return slide

    def _texte(self, slide, texte, x, y, l, h, taille=14, couleur=None, gras=False):
        boite = slide.shapes.add_textbox(x, y, l, h)
        tf = boite.text_frame
        tf.word_wrap = True
        for i, ligne in enumerate(texte.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = ligne
            p.font.size, p.font.bold = Pt(taille), gras
            p.font.color.rgb = couleur or _rgb(self.charte.texte)
        return boite

    # --------------------------------------------------------------- diapositives
    def titre(self, titre: str, lignes: list[str]):
        lignes = [*lignes, f"Généré le {timezone.localtime():%d/%m/%Y à %H:%M}"]
        layout = self.dispositions["titre"]
        if layout is not None:
            slide = self.prs.slides.add_slide(layout)
            ph_titre = self._titre_de(slide)
            ph_titre.text = titre
            autres = [p for p in slide.placeholders if p != ph_titre and p.placeholder_format.type in TYPES_TITRE]
            if autres:
                tf = autres[0].text_frame
                tf.text = "\n".join(lignes)
                for para in tf.paragraphs:  # sans puces, taille fixe
                    para.font.size = Pt(16)
                    ppr = para._p.get_or_add_pPr()
                    ppr.set("marL", "0")
                    ppr.set("indent", "0")
                    ppr.insert(0, etree.SubElement(ppr, "{http://schemas.openxmlformats.org/drawingml/2006/main}buNone"))
            self._vider(slide, garder=[ph_titre, *autres[:1]])
            return
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fond = slide.background.fill
        fond.solid()
        fond.fore_color.rgb = _rgb(self.charte.principale)
        self._texte(slide, titre, Inches(0.8), Inches(2.2), self.largeur - Inches(1.6), Inches(1.4),
                    taille=40, couleur=_rgb(self.charte.secondaire), gras=True)
        self._texte(slide, "\n".join(lignes), Inches(0.8), Inches(3.8), self.largeur - Inches(1.6), Inches(2.5),
                    taille=18, couleur=BLANC)

    def fin(self):
        if self.dispositions["fin"] is not None:
            self._vider(self.prs.slides.add_slide(self.dispositions["fin"]))

    def tableau(self, titre: str, entetes: list[str], lignes: list[list], statuts: list[list[str]] | None = None,
                largeurs: list[float] | None = None):
        """Tableau paginé ; ``statuts`` colore chaque cellule (alerte / critique)."""
        if not lignes:
            slide = self._diapo(titre)
            self._texte(slide, "Aucune ligne.", self.marge, self.haut_contenu, Inches(6), Inches(0.5))
            return
        hauteur_ligne = Inches(0.36)
        par_page = max(4, min(LIGNES_PAR_TABLEAU, int((self.bas_contenu - self.haut_contenu) / hauteur_ligne) - 1))
        pages = [lignes[i:i + par_page] for i in range(0, len(lignes), par_page)]
        for n, page in enumerate(pages):
            slide = self._diapo(titre + (f" ({n + 1}/{len(pages)})" if len(pages) > 1 else ""))
            largeur = self.largeur - 2 * self.marge
            table = slide.shapes.add_table(len(page) + 1, len(entetes), self.marge, self.haut_contenu,
                                           largeur, hauteur_ligne * (len(page) + 1)).table
            if largeurs:
                total = sum(largeurs)
                for j, l in enumerate(largeurs):
                    table.columns[j].width = Emu(int(largeur * l / total))
            for j, e in enumerate(entetes):
                c = table.cell(0, j)
                c.text = e
                c.fill.solid()
                c.fill.fore_color.rgb = _rgb(self.charte.principale)
                para = c.text_frame.paragraphs[0]
                para.font.size, para.font.bold, para.font.color.rgb = Pt(12), True, BLANC
            for i, ligne in enumerate(page, start=1):
                for j, v in enumerate(ligne):
                    c = table.cell(i, j)
                    c.text = str(v)
                    c.text_frame.paragraphs[0].font.size = Pt(11)
                    statut = statuts[n * par_page + i - 1][j] if statuts else ""
                    c.fill.solid()
                    c.fill.fore_color.rgb = FOND_STATUT.get(statut, BLANC)

    def graphique(self, titre: str, categories: list[str], series: list[dict], commentaire: str = ""):
        """``series`` : dicts {nom, valeurs, role?, pointille?} ; ``role`` = principale,
        secondaire ou gris (couleurs de la charte), sinon une couleur par série."""
        slide = self._diapo(titre)
        donnees = CategoryChartData()
        donnees.categories = categories
        for s in series:
            donnees.add_series(s["nom"], [None if v is None or (isinstance(v, float) and math.isnan(v)) else v
                                          for v in s["valeurs"]])
        bas = self.bas_contenu - (Inches(0.8) if commentaire else 0)
        graphe = slide.shapes.add_chart(XL_CHART_TYPE.LINE, self.marge, self.haut_contenu,
                                        self.largeur - 2 * self.marge, bas - self.haut_contenu, donnees).chart
        graphe.has_legend = True
        graphe.legend.position = XL_LEGEND_POSITION.BOTTOM
        graphe.legend.include_in_layout = False
        graphe.legend.font.size = Pt(11)
        graphe.category_axis.tick_labels.font.size = Pt(10)
        graphe.value_axis.tick_labels.font.size = Pt(10)
        for i, (s, serie) in enumerate(zip(series, graphe.series)):
            serie.smooth = False
            ligne = serie.format.line
            couleur = getattr(self.charte, s["role"]) if s.get("role") else self.charte.series[i % len(self.charte.series)]
            ligne.color.rgb = _rgb(couleur)
            ligne.width = Pt(1.25 if s.get("pointille") else 2.25)
            if s.get("pointille"):
                ligne.dash_style = MSO_LINE_DASH_STYLE.DASH
        if commentaire:
            self._texte(slide, commentaire, self.marge, bas + Inches(0.1),
                        self.largeur - 2 * self.marge, Inches(0.7), taille=14)

    def texte(self, titre: str, paragraphes: list[str]):
        slide = self._diapo(titre)
        taille = 16 if len(paragraphes) <= 12 and sum(map(len, paragraphes)) < 900 else 12
        self._texte(slide, "\n".join(paragraphes), self.marge, self.haut_contenu,
                    self.largeur - 2 * self.marge, self.bas_contenu - self.haut_contenu, taille=taille)

    def octets(self) -> bytes:
        self.fin()
        flux = BytesIO()
        self.prs.save(flux)
        return flux.getvalue()
