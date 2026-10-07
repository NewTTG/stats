"""Contenu des rapports : PowerPoint d'une requête ou d'un événement, Excel d'un événement."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

from apps.evenements.analyse import NIVEAUX, Analyse
from apps.kpi.export_excel import REMPLISSAGE, _entete, _largeurs, description_requete
from apps.kpi.service import Resultat
from apps.kpi.views import donnees_graphiques

from .powerpoint import Deck, nombre, pourcentage

LIBELLES_REGLE = {"seuil": "Seuil dépassé", "ecart": "Écart à la référence",
                  "sans_donnees": "Données absentes", "saturation": "Saturation"}
MAX_CELLULES_ANNEXE = 300


def _statut(s: str) -> str:
    return s.capitalize() if s else "OK"


def _annexe(deck: Deck, par_techno):
    lignes = []
    for res in par_techno:
        if res.cellules_sans_donnees:
            liste = res.cellules_sans_donnees[:MAX_CELLULES_ANNEXE]
            suite = f" … (+{len(res.cellules_sans_donnees) - len(liste)})" if len(liste) < len(res.cellules_sans_donnees) else ""
            lignes.append(f"{res.techno} — {len(res.cellules_sans_donnees)} cellule(s) sans aucune donnée :")
            lignes.append(", ".join(liste) + suite)
            lignes.append("")
    deck.texte("Annexe — cellules sans données", lignes or ["Toutes les cellules du périmètre ont des données."])


# ------------------------------------------------------------------ requête KPI

def pptx_requete(resultat: Resultat) -> bytes:
    deck = Deck()
    deck.titre("Statistiques KPI réseau", [f"{cle} : {val}" for cle, val in description_requete(resultat)[:-1]])
    gran = resultat.requete.granularite_temps
    for res in resultat.par_techno:
        lignes, statuts = [], []
        for k in res.kpis:
            v = res.synthese.get(k.code)
            statut = k.statut(None if v != v else v)  # NaN != NaN
            lignes.append([k.libelle + (" ≈" if k.qualite == "approx" else ""), nombre(v), k.unite, _statut(statut)])
            statuts.append(["", statut, "", statut])
        deck.tableau(f"{res.techno} — synthèse sur la période ({res.cellules_demandees} cellules)",
                     ["KPI", "Valeur", "Unité", "Statut"], lignes, statuts, largeurs=[5, 2, 1.2, 1.5])
        graphes = donnees_graphiques(res, gran)
        for k, g in zip(res.kpis, graphes["kpis"]):
            v = res.synthese.get(k.code)
            commentaire = f"{k.libelle} : {nombre(v)} {k.unite} sur l'ensemble de la période et du périmètre."
            if graphes["tronque"]:
                commentaire += " Graphique limité aux 12 premières entités."
            deck.graphique(f"{res.techno} — {g['titre']}", graphes["periodes"],
                           [{"nom": s["nom"], "valeurs": s["valeurs"]} for s in g["series"]], commentaire)
    if resultat.avertissements:
        deck.texte("Avertissements", resultat.avertissements)
    _annexe(deck, resultat.par_techno)
    return deck.octets()


# ------------------------------------------------------------------ événement

def commentaire_kpi(ligne) -> str:
    k = ligne.kpi
    texte = f"{k.libelle} : {nombre(ligne.valeur)} {k.unite} pendant l'événement"
    if ligne.reference is not None:
        texte += f", {pourcentage(ligne.ecart_pct)} par rapport à la référence ({nombre(ligne.reference)} {k.unite})"
    texte += "."
    if ligne.statut:
        texte += f" Statut : {ligne.statut}."
    return texte


def _lignes_anomalies(analyse: Analyse):
    return [[a.gravite.capitalize(), a.techno, a.entite, LIBELLES_REGLE[a.regle], a.message] for a in analyse.anomalies]


def pptx_evenement(analyse: Analyse) -> bytes:
    e, niveau = analyse.evenement, NIVEAUX[analyse.niveau]
    deck = Deck()
    creneaux = [f"{d:%d/%m/%Y %H:%M} → {f:%d/%m/%Y %H:%M}" for d, f in analyse.creneaux]
    if len(creneaux) > 4:
        creneaux = [*creneaux[:4], f"… et {len(creneaux) - 4} autre(s) créneau(x)"]
    deck.titre(e.nom, [
        *([e.type] if e.type else []),
        "Créneaux : " + " ; ".join(creneaux),
        f"Référence : mêmes jours et heures des {analyse.semaines} semaines précédentes",
        f"Analyse par {niveau.lower()} — " + ", ".join(f"{r.nb_cellules} cellules {r.techno}" for r in analyse.par_techno),
    ])
    for res in analyse.par_techno:
        lignes, statuts = [], []
        for s in res.synthese:
            k = s.kpi
            lignes.append([k.libelle + (" ≈" if k.qualite == "approx" else ""), nombre(s.valeur), nombre(s.reference),
                           k.unite, pourcentage(s.ecart_pct), _statut(s.statut)])
            statuts.append(["", s.statut, "", "", s.statut, s.statut])
        deck.tableau(f"{res.techno} — synthèse événement / référence", ["KPI", "Événement", "Référence", "Unité",
                     "Écart", "Statut"], lignes, statuts, largeurs=[5, 1.6, 1.6, 1, 1.3, 1.3])
        for courbe, ligne in zip(res.courbes, res.synthese):
            deck.graphique(f"{res.techno} — {courbe.kpi.libelle} ({courbe.kpi.unite})", res.instants, [
                {"nom": "Événement", "valeurs": courbe.evenement, "couleur": "14284B"},
                {"nom": "Référence", "valeurs": courbe.reference, "couleur": "F0A500"},
                {"nom": "Réf. min", "valeurs": courbe.ref_min, "couleur": "8A94A6", "pointille": True},
                {"nom": "Réf. max", "valeurs": courbe.ref_max, "couleur": "8A94A6", "pointille": True},
            ], commentaire_kpi(ligne))
    deck.tableau(f"Anomalies détectées ({len(analyse.anomalies)})", ["Gravité", "Techno", niveau, "Règle", "Détail"],
                 _lignes_anomalies(analyse),
                 [[a.gravite, "", "", "", ""] for a in analyse.anomalies], largeurs=[1.2, 1, 1.8, 2, 7])
    deck.tableau(f"{niveau}s les plus dégradés", ["Rang", "Techno", niveau, "Critiques", "Alertes", "Concerne"],
                 [[i, c["techno"], c["entite"], c["critiques"], c["alertes"], ", ".join(c["kpis"])]
                  for i, c in enumerate(analyse.classement, start=1)], largeurs=[0.8, 1, 2, 1.2, 1.2, 6])
    if analyse.avertissements:
        deck.texte("Avertissements", analyse.avertissements)
    _annexe(deck, analyse.par_techno)
    return deck.octets()


def xlsx_evenement(analyse: Analyse) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    ws["A1"] = analyse.evenement.nom
    ws["A1"].font = Font(bold=True, size=14, color="14284B")
    ws["A2"] = (f"Référence : {analyse.semaines} semaines précédentes — analyse par {analyse.niveau} — créneaux : "
                + " ; ".join(f"{d:%d/%m/%Y %H:%M} → {f:%d/%m/%Y %H:%M}" for d, f in analyse.creneaux))
    _entete(ws, 4, ["Techno", "KPI", "Unité", "Événement", "Référence", "Écart %", "Statut"])
    ligne = 4
    for res in analyse.par_techno:
        for s in res.synthese:
            ligne += 1
            for col, v in enumerate([res.techno, s.kpi.libelle, s.kpi.unite, s.valeur, s.reference,
                                     s.ecart_pct, _statut(s.statut)], start=1):
                c = ws.cell(row=ligne, column=col, value=v)
                if col in (4, 5, 6):
                    c.number_format = "0.00"
            if s.statut:
                ws.cell(row=ligne, column=4).fill = ws.cell(row=ligne, column=7).fill = REMPLISSAGE[s.statut]
    _largeurs(ws, [10, 40, 8, 12, 12, 10, 10])

    ws = wb.create_sheet("Anomalies")
    _entete(ws, 1, ["Gravité", "Techno", NIVEAUX[analyse.niveau], "Règle", "KPI", "Valeur", "Référence", "Détail"])
    for i, a in enumerate(analyse.anomalies, start=2):
        for col, v in enumerate([a.gravite.capitalize(), a.techno, a.entite, LIBELLES_REGLE[a.regle],
                                 a.kpi.libelle if a.kpi else "", a.valeur, a.reference, a.message], start=1):
            ws.cell(row=i, column=col, value=v)
        ws.cell(row=i, column=1).fill = REMPLISSAGE[a.gravite]
    _largeurs(ws, [10, 8, 18, 20, 30, 10, 10, 70])

    ws = wb.create_sheet("Classement")
    _entete(ws, 1, ["Rang", "Techno", NIVEAUX[analyse.niveau], "Critiques", "Alertes", "Score", "Concerne"])
    for i, c in enumerate(analyse.classement, start=1):
        for col, v in enumerate([i, c["techno"], c["entite"], c["critiques"], c["alertes"], c["score"],
                                 ", ".join(c["kpis"])], start=1):
            ws.cell(row=i + 1, column=col, value=v)
    _largeurs(ws, [6, 8, 18, 10, 10, 8, 60])

    for res in analyse.par_techno:
        ws = wb.create_sheet(f"Courbes {res.techno}")
        titres = ["Heure"]
        for c in res.courbes:
            titres += [f"{c.kpi.libelle} — {x}" for x in ("événement", "référence", "réf. min", "réf. max")]
        _entete(ws, 1, titres)
        for i, instant in enumerate(res.instants):
            valeurs = [instant]
            for c in res.courbes:
                valeurs += [c.evenement[i], c.reference[i], c.ref_min[i], c.ref_max[i]]
            for col, v in enumerate(valeurs, start=1):
                ws.cell(row=i + 2, column=col, value=v)
        _largeurs(ws, [12, *([16] * (len(titres) - 1))])
    flux = BytesIO()
    wb.save(flux)
    return flux.getvalue()
