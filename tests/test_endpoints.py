"""Contrôle transversal des routes : accès, HTTP, erreurs, fichiers et administration."""
from datetime import timedelta
from io import BytesIO
from unittest.mock import Mock

import pytest
from django.contrib import admin
from django.contrib.auth.models import User
from django.core.files.base import ContentFile
from django.test import Client
from django.urls import get_resolver, reverse, URLResolver
from django.utils import timezone
from sqlalchemy.exc import OperationalError

from apps.comptes.models import JournalAudit, Perimetre
from apps.evenements.models import Creneau, Evenement
from apps.rapports.models import Rapport, RapportPlanifie
from .test_service import analyste, base_kpi, lecteur_paita, referentiel, requete
from .test_vues import PARAMS
from .test_evenements import base_horaire, evenement

pytestmark = pytest.mark.django_db


@pytest.fixture
def utilisateur(client, analyste):
    client.force_login(analyste)
    return client


@pytest.mark.parametrize('route', ['recherche', 'evenement'])
def test_source_indisponible_reponse_controlee(utilisateur, referentiel, monkeypatch, route):
    panne = Mock(side_effect=OperationalError('SELECT secret_interne', {}, Exception('serveur-prive')))
    if route == 'recherche':
        monkeypatch.setattr('apps.kpi.views.moteur_kpi', panne)
        url, params = '/', PARAMS
    else:
        e = Evenement.objects.create(nom='Événement accessible', cellules=['AAAe1'])
        monkeypatch.setattr('apps.evenements.views.moteur_kpi', panne)
        url, params = f'/evenements/{e.pk}/', {'debut': '2026-09-01T18:00', 'fin': '2026-09-01T20:00'}
    reponse = utilisateur.get(url, params)
    assert reponse.status_code == 503
    html = reponse.content.decode()
    assert 'momentanément indisponibles' in html
    assert 'serveur-prive' not in html and 'secret_interne' not in html


def test_evenements_hors_perimetre_invisibles(client, referentiel, lecteur_paita, monkeypatch):
    client.force_login(lecteur_paita)
    prive = Evenement.objects.create(nom='SECRET NOUMEA', cellules=['AAAe1'], description='CONFIDENTIEL')
    visible = Evenement.objects.create(nom='PUBLIC PAITA', sites=['BBB200'])
    for e in (prive, visible):
        Creneau.objects.create(evenement=e, debut=timezone.now(), fin=timezone.now() + timedelta(hours=1))
    assert 'SECRET NOUMEA' not in client.get('/evenements/').content.decode()
    assert 'PUBLIC PAITA' in client.get('/evenements/').content.decode()
    assert 'SECRET NOUMEA' not in str(client.get('/suggestions/', {'q': 'SECRET'}).json())
    source = Mock(side_effect=AssertionError('ne doit pas accéder à la source'))
    monkeypatch.setattr('apps.evenements.views.moteur_kpi', source)
    assert client.get(f'/evenements/{prive.pk}/').status_code == 404
    assert client.post(f'/rapports/evenement/{prive.pk}/', {'format': 'pdf'}).status_code == 404
    assert not Rapport.objects.exists()


def test_accueil_kpi_autorise_mais_aucune_cellule(client, base_kpi, monkeypatch):
    user = User.objects.create_user('sans-cellules')
    p = Perimetre.objects.create(nom='Cellules retirées', cellules=['INCONNUE'], kpis_autorises=['lte_rrc_setup_sr'])
    p.utilisateurs.add(user)
    client.force_login(user)
    monkeypatch.setattr('apps.kpi.dashboard.moteur_kpi', lambda: base_kpi)
    reponse = client.get('/dashboard/')
    assert reponse.status_code == 200
    assert reponse.context['dashboard']['cartes'] == []
    assert reponse.context['dashboard']['message']


@pytest.mark.parametrize('requete_invalide', [{}, [], {'techno': ['LTE']}])
def test_lancement_rapport_enregistre_invalide(utilisateur, analyste, requete_invalide):
    plan = RapportPlanifie.objects.create(utilisateur=analyste, titre='Rapport ancien', requete=requete_invalide)
    reponse = utilisateur.post(f'/rapports/enregistres/{plan.pk}/lancer/')
    assert reponse.status_code == 302 and reponse.url == '/rapports/'
    assert not Rapport.objects.exists()
    assert 'invalide' in utilisateur.get('/rapports/').content.decode().lower()


def test_file_indisponible_ne_laisse_pas_un_rapport_en_attente(utilisateur, referentiel, monkeypatch):
    monkeypatch.setattr('apps.rapports.lancement.async_task', Mock(side_effect=RuntimeError('secret serveur')))
    reponse = utilisateur.post('/rapports/requete/', {**PARAMS, 'format': 'pdf'})
    assert reponse.status_code == 302
    rapport = Rapport.objects.get()
    assert rapport.statut == 'erreur'
    assert 'secret serveur' not in utilisateur.get('/rapports/').content.decode()


def test_erreur_generation_ne_divulgue_pas_les_details(utilisateur, analyste, monkeypatch):
    from apps.rapports.taches import generer
    rapport = Rapport.objects.create(utilisateur=analyste, titre='Rapport', nature='requete', format='pdf')
    monkeypatch.setattr('apps.rapports.taches.contenu', Mock(side_effect=RuntimeError('secret serveur SQL')))
    generer(rapport.pk)
    assert 'secret serveur SQL' not in utilisateur.get('/rapports/').content.decode()


def test_admin_audit_malforme_ne_provoque_pas_500(client):
    user = User.objects.create_superuser('admin-audit')
    client.force_login(user)
    JournalAudit.objects.create(utilisateur=user, action='recherche_kpi', requete={'q': 42})
    assert client.get(reverse('admin:comptes_journalaudit_changelist')).status_code == 200


APPLICATION = [
    ('kpi:dashboard', {}, 200),
    ('kpi:requete', {}, 200), ('kpi:suggestions', {}, 200),
    ('evenements:liste', {}, 200), ('evenements:detail', {'pk': 1}, 200),
    # Création / modification : droits evenements.* requis (l'analyste ne les a pas).
    ('evenements:nouveau', {}, 403), ('evenements:modifier', {'pk': 1}, 403), ('evenements:reseau', {}, 403),
    ('evenements:supprimer', {'pk': 1}, 405), ('evenements:memoriser_periode', {'pk': 1}, 405),
    # Réseau : consultation ouverte à qui voit tout le réseau ; déclaration / modification : droits referentiel.*.
    ('reseau:liste', {}, 200), ('reseau:fiche', {'code': 'AAA100'}, 200),
    ('reseau:nouveau', {}, 403), ('reseau:modifier', {'code': 'AAA100'}, 403), ('reseau:supprimer', {'code': 'AAA100'}, 405),
    ('rapports:liste', {}, 200), ('rapports:requete', {}, 405),
    ('rapports:evenement', {'pk': 1}, 405), ('rapports:telecharger', {'pk': 1}, 200),
    ('rapports:enregistrer', {}, 200), ('rapports:planifie_lancer', {'pk': 1}, 405),
    ('rapports:planifie_basculer', {'pk': 1}, 405), ('rapports:planifie_supprimer', {'pk': 1}, 405),
    ('login', {}, 200), ('logout', {}, 405),
]
LECTURES = [nom for nom, _, statut in APPLICATION if statut == 200 and nom not in ('login', 'rapports:enregistrer')]
MUTATIONS = [nom for nom, _, statut in APPLICATION if statut == 405] + ['rapports:enregistrer']


@pytest.fixture
def ressources(analyste, referentiel, base_kpi, settings, tmp_path, monkeypatch):
    settings.MEDIA_ROOT = tmp_path / 'media'
    monkeypatch.setattr('apps.kpi.views.moteur_kpi', lambda: base_kpi)
    monkeypatch.setattr('apps.kpi.dashboard.moteur_kpi', lambda: base_kpi)
    monkeypatch.setattr('apps.evenements.views.moteur_kpi', lambda: base_kpi)
    e = Evenement.objects.create(pk=1, nom='Événement test', cellules=['AAAe1'])
    r = Rapport.objects.create(pk=1, utilisateur=analyste, titre='Fichier test', nature='requete', format='pdf', statut='termine')
    r.fichier.save('test.pdf', ContentFile(b'%PDF-1.4\ncontenu test'))
    p = RapportPlanifie.objects.create(pk=1, utilisateur=analyste, titre='Rapport test',
                                      requete=requete().model_dump(mode='json', exclude={'periode'}))
    return e, r, p


def _url(nom):
    kwargs = next(kwargs for n, kwargs, _ in APPLICATION if n == nom)
    return reverse(nom, kwargs=kwargs)


def test_inventaire_routes_applicatives_complet():
    def noms(patterns, prefix=''):
        for p in patterns:
            if isinstance(p, URLResolver):
                if p.namespace != 'admin':
                    yield from noms(p.url_patterns, prefix + (p.namespace + ':' if p.namespace else ''))
            elif p.name:
                yield prefix + p.name
    assert set(noms(get_resolver().url_patterns)) == {nom for nom, _, _ in APPLICATION}


@pytest.mark.parametrize('nom,kwargs,statut', APPLICATION)
@pytest.mark.parametrize('methode', ['get', 'head'])
def test_routes_get_et_head(utilisateur, ressources, nom, kwargs, statut, methode):
    response = getattr(utilisateur, methode)(reverse(nom, kwargs=kwargs), PARAMS if nom == 'rapports:enregistrer' else {})
    try:
        assert response.status_code == statut
        if nom == 'kpi:suggestions':
            assert response['Content-Type'] == 'application/json'
        if nom == 'rapports:telecharger':
            assert response['Content-Type'] == 'application/pdf'
            assert 'attachment;' in response['Content-Disposition']
    finally:
        response.close()


@pytest.mark.parametrize('nom', LECTURES)
@pytest.mark.parametrize('methode', ['post', 'put', 'patch', 'delete'])
def test_routes_lecture_refusent_les_mutations(utilisateur, ressources, nom, methode):
    assert getattr(utilisateur, methode)(_url(nom)).status_code == 405


@pytest.mark.parametrize('nom', MUTATIONS + ['login'])
def test_csrf_obligatoire(analyste, ressources, nom):
    client = Client(enforce_csrf_checks=True)
    client.force_login(analyste)
    assert client.post(_url(nom), PARAMS).status_code == 403
    assert Rapport.objects.count() == 1 and RapportPlanifie.objects.count() == 1


@pytest.mark.parametrize('nom,kwargs,statut', [a for a in APPLICATION if a[0] not in ('login', 'logout')])
@pytest.mark.parametrize('methode', ['get', 'post'])
def test_connexion_obligatoire_sans_visiteurs(client, settings, nom, kwargs, statut, methode):
    settings.ACCES_ANONYME = False
    settings.LOGIN_URL = 'login'
    reponse = getattr(client, methode)(reverse(nom, kwargs=kwargs))
    assert reponse.status_code == 302
    assert reponse.url.startswith('/connexion/?next=')
    assert not User.objects.exists()


@pytest.mark.parametrize('nom', ['evenements:detail', 'rapports:telecharger', 'rapports:evenement',
                                  'rapports:planifie_lancer', 'rapports:planifie_basculer', 'rapports:planifie_supprimer'])
def test_identifiant_inexistant(utilisateur, nom):
    methode = utilisateur.get if nom in ('evenements:detail', 'rapports:telecharger') else utilisateur.post
    assert methode(reverse(nom, kwargs={'pk': 999999})).status_code == 404


@pytest.mark.parametrize('nom', ['rapports:telecharger', 'rapports:planifie_lancer',
                                  'rapports:planifie_basculer', 'rapports:planifie_supprimer'])
def test_ressources_autre_utilisateur_inaccessibles(client, ressources, nom):
    client.force_login(User.objects.create_superuser('autre-administrateur'))
    methode = client.get if nom == 'rapports:telecharger' else client.post
    assert methode(_url(nom)).status_code == 404
    assert 'Fichier test' not in client.get('/rapports/').content.decode()
    assert 'Rapport test' not in client.get('/rapports/').content.decode()


@pytest.mark.parametrize('statut', ['en_attente', 'en_cours', 'erreur'])
def test_rapport_inacheve_non_telechargeable(utilisateur, ressources, statut):
    Rapport.objects.filter(pk=1).update(statut=statut)
    assert utilisateur.get('/rapports/1/telecharger/').status_code == 404


def test_fichier_disparu_redirection_sans_erreur(utilisateur, ressources):
    ressources[1].fichier.delete(save=False)
    r = utilisateur.get('/rapports/1/telecharger/')
    assert r.status_code == 302 and r.url == '/rapports/'
    assert 'introuvable' in utilisateur.get('/rapports/').content.decode()


@pytest.mark.parametrize('parametres', [
    {'techno': 'GSM'}, {'debut': 'invalide'}, {'fin': '0001-01-01'},
    {'perimetre_type': 'inexistant'}, {'granularite_temps': 'invalide'},
    {'kpis': ['secret']}, {'fenetre_horaire': '24-48'},
])
def test_parametres_recherche_invalides(utilisateur, referentiel, parametres):
    r = utilisateur.get('/', {**PARAMS, **parametres})
    assert r.status_code == 200 and r.context['erreur']
    r = utilisateur.post('/rapports/requete/', {**PARAMS, **parametres})
    assert r.status_code == 302 and not Rapport.objects.exists()


@pytest.mark.parametrize('url', ['/rapports/requete/', '/rapports/evenement/1/'])
def test_format_rapport_invalide(utilisateur, ressources, url):
    assert utilisateur.post(url, {**PARAMS, 'format': 'exe'}).status_code == 404
    assert Rapport.objects.count() == 1


def test_redirection_connexion_externe_refusee(client):
    User.objects.create_user('connexion-test', password='mot-de-passe-test')
    reponse = client.post('/connexion/', {'username': 'connexion-test', 'password': 'mot-de-passe-test',
                                         'next': 'https://autre.example/vol'})
    assert reponse.status_code == 302 and reponse.url == '/'


def test_favicon_et_page_inconnue(client):
    assert client.get('/favicon.ico').status_code == 204
    assert client.get('/route-inconnue/').status_code == 404


MODELES_ADMIN = [(m._meta.app_label, m._meta.model_name) for m in admin.site._registry]


@pytest.mark.parametrize('app,modele', MODELES_ADMIN)
@pytest.mark.parametrize('page', ['changelist', 'add', 'change', 'history', 'delete'])
@pytest.mark.parametrize('role', ['admin', 'lecteur', 'staff_sans_droits'])
def test_routes_modeles_administration(client, ressources, app, modele, page, role):
    user = User.objects.create_user('controle-admin', is_staff=role != 'lecteur', is_superuser=role == 'admin')
    client.force_login(user)
    args = ['1'] if page in ('change', 'history', 'delete') else []
    r = client.get(reverse(f'admin:{app}_{modele}_{page}', args=args))
    if role == 'lecteur':
        assert r.status_code == 302 and '/admin/login/' in r.url
    elif role == 'staff_sans_droits':
        # Django renvoie vers l'index pour l'historique d'un objet inexistant.
        assert r.status_code == 403 or (page == 'history' and r.status_code == 302 and r.url == '/admin/')
    else:
        assert r.status_code in (200, 302, 403)  # ajouts/suppressions volontairement interdits ; objet absent


@pytest.mark.parametrize('nom,kwargs,statut', [
    ('admin:index', {}, 200), ('admin:login', {}, 302), ('admin:logout', {}, 405),
    ('admin:password_change', {}, 200), ('admin:password_change_done', {}, 200),
    ('admin:autocomplete', {}, 403), ('admin:jsi18n', {}, 200),
    ('admin:view_on_site', {'content_type_id': '999999', 'object_id': '1'}, 404),
    ('admin:auth_user_password_change', {'id': '1'}, 200),
])
def test_routes_generales_admin(client, ressources, nom, kwargs, statut):
    client.force_login(User.objects.create_superuser('admin-general'))
    assert client.get(reverse(nom, kwargs=kwargs)).status_code == statut


@pytest.mark.parametrize('app', sorted({app for app, _ in MODELES_ADMIN}))
def test_sections_admin(client, ressources, app):
    client.force_login(User.objects.create_superuser('admin-section'))
    assert client.get(reverse('admin:app_list', kwargs={'app_label': app})).status_code == 200


def test_admin_sans_permission_ne_modifie_pas_les_reglages(client):
    from apps.evenements.models import ReglagesAnomalies
    from apps.kpi.models import SeuilKpi
    client.force_login(User.objects.create_user('staff-vide', is_staff=True))
    SeuilKpi.objects.all().delete()
    ReglagesAnomalies.objects.all().delete()
    assert client.get(reverse('admin:kpi_seuilkpi_changelist')).status_code == 403
    assert client.get(reverse('admin:evenements_reglagesanomalies_changelist')).status_code == 403
    assert not SeuilKpi.objects.exists() and not ReglagesAnomalies.objects.exists()


@pytest.mark.parametrize('format_rapport', ['pdf', 'pptx', 'xlsx'])
@pytest.mark.parametrize('nature', ['requete', 'evenement'])
def test_generation_et_telechargement_tous_formats(utilisateur, analyste, referentiel, evenement,
                                                  base_kpi, base_horaire, settings, tmp_path,
                                                  monkeypatch, nature, format_rapport):
    from apps.rapports.taches import generer
    settings.MEDIA_ROOT = tmp_path / 'rapports'
    monkeypatch.setattr('apps.rapports.taches.moteur_kpi', lambda: base_kpi if nature == 'requete' else base_horaire)
    monkeypatch.setattr('apps.rapports.lancement.async_task', lambda _nom, pk, **kw: generer(pk))
    url = '/rapports/requete/' if nature == 'requete' else f'/rapports/evenement/{evenement.pk}/'
    reponse = utilisateur.post(url, {**PARAMS, 'format': format_rapport, 'niveau': 'secteur'})
    assert reponse.status_code == 302 and reponse.url == '/rapports/'
    rapport = Rapport.objects.get()
    assert rapport.statut == 'termine', rapport.message
    r = utilisateur.get(f'/rapports/{rapport.pk}/telecharger/')
    assert r.status_code == 200
    try:
        contenu = BytesIO(b''.join(r.streaming_content))
    finally:
        r.close()
    if format_rapport == 'pdf':
        from pypdf import PdfReader
        assert PdfReader(contenu).pages
    elif format_rapport == 'xlsx':
        from openpyxl import load_workbook
        assert load_workbook(contenu).sheetnames
    else:
        from pptx import Presentation
        assert len(Presentation(contenu).slides) > 0


@pytest.mark.parametrize('declaration', [{'cellules': ['BBBe1']}, {'sites': ['BBB200']}, {'secteurs': ['BBB2001']}])
def test_evenement_visible_selon_les_trois_rattachements(client, lecteur_paita, referentiel, base_kpi, monkeypatch, declaration):
    client.force_login(lecteur_paita)
    e = Evenement.objects.create(nom='Visible', **declaration)
    monkeypatch.setattr('apps.evenements.views.moteur_kpi', lambda: base_kpi)
    assert client.get(f'/evenements/{e.pk}/').status_code == 200


def test_recherche_evenement_inaccessible_comme_inconnu(client, referentiel, lecteur_paita, base_kpi, monkeypatch):
    client.force_login(lecteur_paita)
    Evenement.objects.create(nom='SECRET', cellules=['AAAe1'])
    monkeypatch.setattr('apps.kpi.views.moteur_kpi', lambda: base_kpi)
    erreurs = []
    for nom in ('SECRET', 'INEXISTANT'):
        r = client.get('/', {**PARAMS, 'perimetre_type': 'evenement', 'perimetre_valeurs': nom})
        assert r.status_code == 200
        erreurs.append(r.context['erreur'].replace(nom, 'X'))
    assert erreurs[0] == erreurs[1]


@pytest.mark.parametrize('app,modele', MODELES_ADMIN)
def test_anciennes_urls_admin_redirigent(client, ressources, app, modele):
    client.force_login(User.objects.create_superuser('admin-ancienne-url'))
    r = client.get(f'/admin/{app}/{modele}/1/')
    assert r.status_code == 302
    assert r.url == reverse(f'admin:{app}_{modele}_change', args=['1'])


def test_routes_inconnues_admin(client):
    client.force_login(User.objects.create_superuser('admin-url-inconnue'))
    assert client.get('/admin/route-inconnue/').status_code == 404


@pytest.mark.parametrize('url', ['/admin/logout/', '/admin/password_change/', '/admin/auth/user/1/delete/'])
def test_csrf_admin_obligatoire(url):
    client = Client(enforce_csrf_checks=True)
    client.force_login(User.objects.create_superuser('admin-csrf'))
    assert client.post(url, {}).status_code == 403


def test_fichiers_statiques_servis_en_production(settings, tmp_path):
    from django.core.management import call_command
    from io import StringIO
    settings.DEBUG = False
    settings.STATIC_ROOT = tmp_path / 'staticfiles'
    settings.WHITENOISE_USE_FINDERS = False
    call_command('collectstatic', interactive=False, verbosity=0, stdout=StringIO())
    client = Client()
    for chemin, mime in [('css/app.css', 'text/css'), ('css/workspace.css', 'text/css'),
                         ('js/kpi.js', 'javascript'), ('vendor/echarts-5.6.0.min.js', 'javascript')]:
        r = client.get('/static/' + chemin)
        try:
            assert r.status_code == 200
            assert mime in r['Content-Type']
            assert b''.join(r.streaming_content)
        finally:
            r.close()
    assert not User.objects.exists()  # les fichiers ne créent pas de visiteurs
