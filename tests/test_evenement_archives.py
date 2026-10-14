"""Régressions des éditions : période indépendante, archives et déplacement de remorque."""
from datetime import datetime
from io import BytesIO

import pandas as pd
import pytest
from django.contrib.auth.models import Permission, User
from django.http import QueryDict
from django.urls import reverse
from openpyxl import load_workbook

from apps.comptes.models import Perimetre
from apps.evenements import archives, periode
from apps.evenements.analyse import analyser
from apps.evenements.models import EditionEvenement, Evenement
from apps.kpi.models import SeuilKpi
from apps.rapports import taches
from apps.rapports.models import Rapport
from apps.referentiel.models import Site

from .test_evenements import base_horaire, _synthese  # noqa: F401
from .test_service import referentiel, analyste  # noqa: F401
from .test_sites_mobiles import remorque  # noqa: F401

pytestmark = pytest.mark.django_db


def plage(annee=2026, debut=18, fin=21):
    return [(datetime(annee, 9, 29, debut), datetime(annee, 9, 29, fin))]


@pytest.fixture
def editions(base_horaire, analyste, remorque):
    e = Evenement.objects.create(nom="Carnaval", sites=["AAA100"],
                                kpis=["lte_payload_dl", "lte_rrc_setup_sr"])
    lignes = pd.read_sql_table("lte_cell_hour", base_horaire)
    ancien = lignes.copy()
    ancien["DateHour"] = pd.to_datetime(ancien["DateHour"]) - pd.DateOffset(years=1)
    ancien.to_sql("lte_cell_hour", base_horaire, if_exists="append", index=False)
    a2025 = archives.enregistrer(e, analyste, base_horaire, plage(2025), "Carnaval 2025")
    mobiles = lignes[lignes["EutranCell_Id"] == "AAAe1"].copy()
    mobiles["EutranCell_Id"] = "1RMe1"
    mobiles["PayloadDl_mB"] = 100
    mobiles.to_sql("lte_cell_hour", base_horaire, if_exists="append", index=False)
    e.deploiements = {"1RM128": {"lat": -22.2, "lon": 166.4, "secteurs": [
        {"numero": 1, "azimut": 0, "cellules": [["1RMe1", "LTE"]]}]}}
    e.save()
    a2026 = archives.enregistrer(e, analyste, base_horaire, plage(), "Carnaval 2026")
    return e, a2025, a2026


def test_deplacement_ne_modifie_pas_les_deux_editions(editions, analyste):
    e, ancien, dernier = editions
    original = ancien.donnees
    # Déplacement, changement de périmètre, nom du site et seuils après l'événement.
    e.deploiements = {}
    e.sites = []
    e.kpis = []
    e.save()
    Site.objects.filter(code_site="AAA100").update(nom="NOUVEL_EMPLACEMENT")
    SeuilKpi.synchroniser()
    SeuilKpi.objects.filter(code="lte_rrc_setup_sr").update(critique=10, alerte=20)
    a = analyser(e, analyste, None, creneaux=plage(), reference=plage(2025),
                 archive=dernier, archive_reference=ancien)
    vol = _synthese(a, "lte_payload_dl")
    assert vol.valeur == 420 and vol.reference == 120
    assert a.par_techno[0].nb_cellules == 3
    assert set(a.par_techno[0].courbes[0].par_site) == {"SITE_AAA", "REMORQUE_1"}
    assert _synthese(a, "lte_rrc_setup_sr").statut == "critique"
    ancien.refresh_from_db()
    assert ancien.donnees == original and "1RMe1" not in ancien.donnees["technos"]["LTE"]["cellules"]
    assert EditionEvenement.objects.count() == 2


def test_reference_plus_longue_garde_toutes_les_heures(editions, analyste):
    e, ancien, dernier = editions
    a = analyser(e, analyste, None, creneaux=plage(fin=20), reference=plage(2025),
                 archive=dernier, archive_reference=ancien)
    vol = _synthese(a, "lte_payload_dl")
    assert vol.valeur == 280 and vol.reference == 120
    courbe = next(c for c in a.par_techno[0].courbes if c.kpi.code == "lte_payload_dl")
    assert courbe.evenement == [140, 140, None] and courbe.reference == [40, 40, 40]
    assert any("2 h" in x and "3 h" in x for x in a.avertissements)


def test_archive_refuse_les_dates_hors_couverture(editions, analyste):
    from apps.evenements.analyse import AnalyseImpossible
    e, ancien, dernier = editions
    with pytest.raises(AnalyseImpossible, match="dépasse"):
        analyser(e, analyste, None, creneaux=plage(fin=22), reference=plage(2025),
                 archive=dernier, archive_reference=ancien)


def test_reference_automatique_et_page_sans_source(client, editions, analyste, monkeypatch):
    e, ancien, dernier = editions
    client.force_login(analyste)
    monkeypatch.setattr("apps.evenements.views.moteur_kpi", lambda: pytest.fail("aucun accès source attendu"))
    r = client.get(reverse("evenements:detail", args=[e.pk]))
    assert r.status_code == 200 and r.context["archive"].pk == dernier.pk
    assert r.context["reference"] == plage(2025)
    assert r.context["origine_reference"] == "Édition enregistrée : Carnaval 2025"
    assert _synthese(r.context["analyse"], "lte_payload_dl").reference == 120
    assert "29/09/2025" in r.content.decode() and EditionEvenement.objects.count() == 2


def test_droits_reappliques_aux_archives(client, editions):
    e, ancien, dernier = editions
    u = User.objects.create_user("restreint")
    Perimetre.objects.create(nom="AAA", sites=["AAA100"], kpis_autorises=["lte_payload_dl"]).utilisateurs.add(u)
    a = analyser(e, u, None, creneaux=plage(), reference=plage(2025), archive=dernier, archive_reference=ancien)
    assert a.par_techno[0].nb_cellules == 2
    assert [s.kpi.code for s in a.par_techno[0].synthese] == ["lte_payload_dl"]
    assert _synthese(a, "lte_payload_dl").valeur == 120
    u.user_permissions.add(Permission.objects.get(codename="change_evenement"))
    client.force_login(u)
    assert client.post(reverse("evenements:enregistrer_edition", args=[e.pk])).status_code == 403
    autre = Evenement.objects.create(nom="Autre", sites=["AAA100"])
    assert client.get(reverse("evenements:detail", args=[autre.pk]), {"edition_id": dernier.pk}).status_code == 404


def test_reference_calendaire_exclut_remorque(editions, base_horaire, analyste):
    e, ancien, dernier = editions
    # Même si la remorque avait du trafic ailleurs en 2025, il ne fait pas partie de la référence.
    lignes = pd.read_sql_table("lte_cell_hour", base_horaire)
    mobiles = lignes[lignes["EutranCell_Id"] == "1RMe1"].copy()
    mobiles["DateHour"] = pd.to_datetime(mobiles["DateHour"]) - pd.DateOffset(years=1)
    mobiles.to_sql("lte_cell_hour", base_horaire, if_exists="append", index=False)
    a = analyser(e, analyste, base_horaire, creneaux=plage(), reference=plage(2025))
    assert _synthese(a, "lte_payload_dl").reference == 120


def test_formulaire_personnalise_et_exports(client, editions, analyste, monkeypatch):
    e, ancien, dernier = editions
    client.force_login(analyste)
    monkeypatch.setattr("apps.evenements.views.moteur_kpi", lambda: None)
    # Dates invalides : aucune analyse n'est lancée, saisie conservée.
    r = client.get(reverse("evenements:detail", args=[e.pk]), {
        "edition_id": dernier.pk, "comparaison": "personnalisee", "ref_detail": "1",
        "ref_jour": "2025-09-29", "ref_de": "18", "ref_a": ""})
    assert "analyse" not in r.context and 'value="2025-09-29"' in r.content.decode()
    monkeypatch.setattr("apps.rapports.views.lancer", lambda rapport, parent=None: True)
    client.post(reverse("rapports:evenement", args=[e.pk]), {
        "edition_id": dernier.pk, "comparaison": str(ancien.pk), "du": "2026-09-29", "de": "18", "a": "21", "format": "xlsx"})
    rapport = Rapport.objects.get()
    assert rapport.parametres["archive_reference"] == ancien.pk
    monkeypatch.setattr(taches, "moteur_kpi", lambda: pytest.fail("export archivé autonome"))
    wb = load_workbook(BytesIO(taches.contenu(rapport)))
    assert "29/09/2025" in wb.active["A2"].value
    assert any(row[3:5] == (420, 120) for row in wb.active.iter_rows(min_row=5, values_only=True))


def test_enregistrer_edition_post_uniquement(client, editions, analyste, base_horaire, monkeypatch):
    e, ancien, dernier = editions
    analyste.user_permissions.add(Permission.objects.get(codename="change_evenement"))
    client.force_login(analyste)
    monkeypatch.setattr("apps.evenements.views.moteur_kpi", lambda: base_horaire)
    url = reverse("evenements:enregistrer_edition", args=[e.pk])
    assert client.get(url).status_code == 405
    data = {"du": "2026-09-29", "de": "18", "a": "21", "nom_edition": "Carnaval corrigé"}
    r = client.post(url, data)
    assert r.status_code == 302 and "edition_id=" in r["Location"]
    client.post(url, data)
    assert EditionEvenement.objects.count() == 3  # pas d'écrasement
    client.post(url, {**data, "du": "2099-01-01", "nom_edition": "Futur"})
    assert EditionEvenement.objects.count() == 3


def test_pas_de_double_comptage_des_creneaux():
    with pytest.raises(periode.PeriodeInvalide, match="chevaucher"):
        periode.lire(QueryDict("detail=1&jour=2026-09-29&de=18&a=21&jour=2026-09-29&de=20&a=22"))


def test_edition_reste_consultable_sans_source(client, editions, analyste, monkeypatch):
    from apps.kpi.source import BaseKpiNonConfiguree
    e, ancien, dernier = editions
    client.force_login(analyste)
    def indisponible():
        raise BaseKpiNonConfiguree("indisponible")
    monkeypatch.setattr("apps.evenements.views.moteur_kpi", indisponible)
    r = client.get(reverse("evenements:detail", args=[e.pk]), {"edition_id": dernier.pk, "comparaison": "calendrier"})
    assert r.status_code == 200
    assert _synthese(r.context["analyse"], "lte_payload_dl").valeur == 420
    assert _synthese(r.context["analyse"], "lte_payload_dl").reference is None
    assert any("référence en N/D" in a for a in r.context["analyse"].avertissements)


def test_composites_conserves():
    from apps.kpi.catalogue import DefinitionKpi, catalogue
    cat = catalogue()
    composites = [k for k in cat.values() if k.produit_de]
    assert composites
    for k in composites:
        relu = DefinitionKpi(**archives.definition(k))
        assert relu.composants == k.composants


def test_un_kpi_nouveau_reste_visible_sans_reference(editions, analyste):
    e, ancien, dernier = editions
    ancien.donnees["catalogue"].pop("lte_rrc_setup_sr")
    a = analyser(e, analyste, None, creneaux=plage(), reference=plage(2025), archive=dernier, archive_reference=ancien)
    assert _synthese(a, "lte_rrc_setup_sr").valeur is not None
    assert _synthese(a, "lte_rrc_setup_sr").reference is None


def test_compteurs_nuls_restent_en_nd(editions, analyste):
    e, ancien, dernier = editions
    for ligne in ancien.donnees["technos"]["LTE"]["lignes"]:
        ligne["pmRrcConnEstabSucc"] = None
    a = analyser(e, analyste, None, creneaux=plage(), reference=plage(2025), archive=dernier, archive_reference=ancien)
    assert _synthese(a, "lte_rrc_setup_sr").reference is None
    assert _synthese(a, "lte_payload_dl").reference == 120


def test_comparaison_personnalisee_de_bout_en_bout(client, editions, analyste, base_horaire, monkeypatch):
    e, ancien, dernier = editions
    client.force_login(analyste)
    monkeypatch.setattr("apps.evenements.views.moteur_kpi", lambda: base_horaire)
    r = client.get(reverse("evenements:detail", args=[e.pk]), {
        "edition_id": "direct", "du": "2026-09-29", "de": "18", "a": "22",
        "comparaison": "personnalisee", "ref_detail": "1", "ref_jour": "2025-09-29", "ref_de": "18", "ref_a": "21"})
    assert r.context["reference"] == plage(2025)
    vol = _synthese(r.context["analyse"], "lte_payload_dl")
    assert vol.valeur == 560 and vol.reference == 120
    assert any("4 h" in a and "3 h" in a for a in r.context["analyse"].avertissements)
    assert "ref_jour=2025-09-29" in r.content.decode()  # onglets de niveau préservent la référence


def test_formulaire_dans_un_navigateur(client, editions, analyste, monkeypatch):
    import shutil
    from pathlib import Path
    pw = pytest.importorskip("playwright.sync_api")
    executable = shutil.which("chromium")
    if not executable:
        pytest.skip("Chromium absent")
    e, ancien, dernier = editions
    client.force_login(analyste)
    monkeypatch.setattr("apps.evenements.views.moteur_kpi", lambda: None)
    r = client.get(reverse("evenements:detail", args=[e.pk]))
    racine = Path(__file__).resolve().parents[1]
    with pw.sync_playwright() as p:
        browser = p.chromium.launch(executable_path=executable, args=["--no-sandbox"])
        page = browser.new_page(viewport={"width": 1440, "height": 1000})
        erreurs = []
        page.on("pageerror", lambda e: erreurs.append(str(e)))
        page.route("**/*", lambda route: route.abort())
        page.set_content(r.content.decode())
        for nom in ("app", "workspace", "evenement-periode"):
            page.add_style_tag(path=str(racine / f"static/css/{nom}.css"))
        page.add_style_tag(path=str(racine / "static/vendor/leaflet-1.9.4/leaflet.css"))
        for nom in ("evenement-periode", "evenement-comparaison"):
            page.add_script_tag(path=str(racine / f"static/js/{nom}.js"))
        for chemin in ("vendor/leaflet-1.9.4/leaflet.js", "js/evenement-zone.js",
                       "vendor/echarts-5.6.0.min.js", "js/kpi.js"):
            page.add_script_tag(path=str(racine / "static" / chemin))
        assert page.locator(".graphiques.courbes canvas").count() >= 1
        assert page.locator('[name="ref_jour"]').is_disabled()
        page.locator('[data-comparaison]').select_option("personnalisee")
        assert page.locator('[name="ref_jour"]').is_visible()
        page.locator('[name="ref_jour"]').fill("2025-08-20")
        page.locator('[name="ref_de"]').select_option("18")
        page.locator('[name="ref_a"]').select_option("21")
        page.locator('[data-reference-ajouter]').click()
        assert page.locator('[data-reference-ligne]').count() == 2
        page.locator('[data-reference-retirer]').last.click()
        valeurs = page.locator('form[data-periode]').evaluate("e => [...new FormData(e).entries()]")
        assert ["ref_jour", "2025-08-20"] in valeurs and ["ref_a", "21"] in valeurs
        assert ["du", "2026-09-29"] in valeurs and ["a", "21"] in valeurs
        (racine / "work").mkdir(exist_ok=True)
        page.screenshot(path=str(racine / "work/evenement-desktop.png"), full_page=True)
        page.set_viewport_size({"width": 390, "height": 844})
        page.wait_for_function("document.documentElement.scrollWidth <= innerWidth", timeout=5000)
        page.screenshot(path=str(racine / "work/evenement-mobile.png"), full_page=True)
        assert not erreurs
        browser.close()
