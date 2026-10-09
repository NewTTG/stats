"""Rapports enregistrés : période relative, échéances, lancements planifiés, écrans et droits."""

from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from django.contrib.auth.models import User
from django.core.management import call_command
from django_q.models import Schedule
from openpyxl import load_workbook

from apps.kpi.forms import champs_depuis_requete
from apps.rapports.models import Rapport, RapportPlanifie
from apps.rapports.planification import lancer_dus, requete_du

from .test_service import analyste, base_kpi, referentiel, requete  # noqa: F401  (fixtures)

pytestmark = pytest.mark.django_db
NOUMEA = ZoneInfo("Pacific/Noumea")


def _a(*args) -> datetime:
    return datetime(*args, tzinfo=NOUMEA)


def _plan(user, **modifs) -> RapportPlanifie:
    valeurs = {"utilisateur": user, "titre": "Hebdo Nouméa", "periode": "hier", "format": "xlsx",
               "requete": requete().model_dump(mode="json", exclude={"periode"}), **modifs}
    plan = RapportPlanifie(**valeurs)
    plan.save()
    return plan


@pytest.mark.parametrize("frequence,jour_semaine,apres,attendu", [
    ("quotidienne", 0, _a(2026, 10, 9, 6, 0), _a(2026, 10, 9, 7, 0)),
    ("quotidienne", 0, _a(2026, 10, 9, 7, 0), _a(2026, 10, 10, 7, 0)),   # strictement après
    ("hebdomadaire", 0, _a(2026, 10, 9, 12, 0), _a(2026, 10, 12, 7, 0)),  # vendredi -> lundi
    ("hebdomadaire", 4, _a(2026, 10, 9, 6, 59), _a(2026, 10, 9, 7, 0)),
    ("mensuelle", 0, _a(2026, 10, 9, 12, 0), _a(2026, 11, 1, 7, 0)),
    ("mensuelle", 0, _a(2026, 12, 15, 12, 0), _a(2027, 1, 1, 7, 0)),
    ("aucune", 0, _a(2026, 10, 9, 12, 0), None),
])
def test_prochain_lancement(analyste, frequence, jour_semaine, apres, attendu):  # noqa: F811
    plan = RapportPlanifie(utilisateur=analyste, frequence=frequence, jour_semaine=jour_semaine, heure=7)
    assert plan.suivant(apres) == attendu


def test_planifier_inactif(analyste):  # noqa: F811
    plan = RapportPlanifie(utilisateur=analyste, frequence="quotidienne", actif=False)
    plan.planifier(_a(2026, 10, 9, 12, 0))
    assert plan.prochain_lancement is None


def test_requete_du(analyste):  # noqa: F811
    req, libelle = requete_du(_plan(analyste, periode="semaine_derniere"), date(2026, 10, 9))
    assert (req.periode.debut, req.periode.fin) == (date(2026, 9, 28), date(2026, 10, 4))
    assert req.kpis == requete().kpis
    assert libelle == "Semaine dernière, 28/09/2026-04/10/2026"


@pytest.fixture
def worker(monkeypatch, base_kpi, settings, tmp_path):  # noqa: F811
    """Exécute la génération immédiatement (pas de worker django-q2 en test)."""
    from apps.rapports import taches

    settings.MEDIA_ROOT = tmp_path
    monkeypatch.setattr("apps.rapports.taches.moteur_kpi", lambda: base_kpi)
    monkeypatch.setattr("apps.rapports.lancement.async_task", lambda _nom, pk, **kw: taches.generer(pk))


def test_lancer_dus(referentiel, analyste, worker):  # noqa: F811
    maintenant = _a(2026, 9, 3, 7, 5)
    du = _plan(analyste, frequence="quotidienne", prochain_lancement=_a(2026, 9, 3, 7, 0))
    futur = _plan(analyste, titre="Futur", frequence="quotidienne", prochain_lancement=_a(2026, 9, 4, 7, 0))
    suspendu = _plan(analyste, titre="Suspendu", frequence="quotidienne", actif=False,
                     prochain_lancement=_a(2026, 9, 1, 7, 0))

    assert lancer_dus(maintenant) == 1
    rapport = Rapport.objects.get()
    assert (rapport.planification, rapport.statut, rapport.format) == (du, "termine", "xlsx")
    assert rapport.titre == "Hebdo Nouméa — Hier, 02/09/2026-02/09/2026"
    assert rapport.parametres["requete"]["periode"] == {"debut": "2026-09-02", "fin": "2026-09-02"}
    wb = load_workbook(BytesIO(rapport.fichier.read()))
    assert wb.sheetnames
    du.refresh_from_db()
    assert (du.prochain_lancement, du.dernier_lancement) == (_a(2026, 9, 4, 7, 0), maintenant)
    futur.refresh_from_db(), suspendu.refresh_from_db()
    assert futur.dernier_lancement is None and suspendu.dernier_lancement is None
    # échéance déjà passée : pas de second lancement
    assert lancer_dus(maintenant) == 0


def test_tache_periodique_creee():
    tache = Schedule.objects.get(name="rapports-planifies")
    assert (tache.func, tache.schedule_type, tache.minutes) == ("apps.rapports.planification.lancer_dus", "I", 15)


def test_commande(analyste, worker, capsys):  # noqa: F811
    call_command("lancer_rapports_planifies")
    assert "0 rapport(s) lancé(s)." in capsys.readouterr().out


@pytest.fixture
def client_analyste(client, analyste):  # noqa: F811
    client.force_login(analyste)
    return client


def _champs_requete():
    return champs_depuis_requete(requete(periode={"debut": date(2026, 9, 1), "fin": date(2026, 9, 2)}))


def test_ecran_enregistrer(client_analyste, referentiel):  # noqa: F811
    html = client_analyste.get("/rapports/enregistrer/", _champs_requete()).content.decode()
    assert "Enregistrer un rapport" in html
    assert 'value="KPI LTE — réseau"' in html


def test_enregistrer_puis_lancer(client_analyste, analyste, referentiel, worker):  # noqa: F811
    donnees = {**_champs_requete(), "titre": "Mensuel réseau", "periode": "mois_dernier", "format": "pptx",
               "frequence": "mensuelle", "jour_semaine": 0, "heure": 6}
    reponse = client_analyste.post("/rapports/enregistrer/", donnees)
    assert reponse.status_code == 302
    plan = RapportPlanifie.objects.get()
    assert (plan.utilisateur, plan.frequence, plan.heure) == (analyste, "mensuelle", 6)
    assert "periode" not in plan.requete and plan.requete["kpis"] == requete().kpis
    assert plan.prochain_lancement.astimezone(NOUMEA).day == 1

    html = client_analyste.get("/rapports/").content.decode()
    assert "Mensuel réseau" in html and "Le 1er du mois à 6 h" in html

    client_analyste.post(f"/rapports/enregistres/{plan.pk}/lancer/")
    assert Rapport.objects.get().planification == plan

    client_analyste.post(f"/rapports/enregistres/{plan.pk}/basculer/")
    plan.refresh_from_db()
    assert not plan.actif and plan.prochain_lancement is None
    client_analyste.post(f"/rapports/enregistres/{plan.pk}/supprimer/")
    assert not RapportPlanifie.objects.exists() and Rapport.objects.count() == 1


def test_requete_invalide(client_analyste):
    reponse = client_analyste.get("/rapports/enregistrer/", {"techno": "LTE"})
    assert reponse.status_code == 302 and reponse.url == "/"


def test_rapport_d_un_autre_utilisateur(client, analyste):  # noqa: F811
    plan = _plan(analyste)
    autre = User.objects.create_user("autre")
    client.force_login(autre)
    for action in ("lancer", "basculer", "supprimer"):
        assert client.post(f"/rapports/enregistres/{plan.pk}/{action}/").status_code == 404
    assert "Hebdo Nouméa" not in client.get("/rapports/").content.decode()


def test_un_echec_ne_bloque_pas_les_autres(referentiel, analyste, worker):  # noqa: F811
    echeance = _a(2026, 9, 3, 7, 0)
    casse = _plan(analyste, titre="Cassé", frequence="quotidienne", prochain_lancement=echeance,
                  requete={"techno": []})  # requête devenue invalide
    bon = _plan(analyste, frequence="quotidienne", prochain_lancement=echeance)
    assert lancer_dus(_a(2026, 9, 3, 7, 5)) == 1
    assert Rapport.objects.get().planification == bon
    casse.refresh_from_db()
    assert casse.prochain_lancement == _a(2026, 9, 4, 7, 0)  # pas relancé en boucle
