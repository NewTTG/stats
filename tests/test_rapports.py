"""Rapports : contenu PowerPoint / Excel, tâche de fond, historique, droits ; écrans événements."""

from io import BytesIO

import pytest
from django.contrib.auth.models import User
from openpyxl import load_workbook
from pptx import Presentation

from apps.comptes.models import JournalAudit
from apps.evenements.analyse import analyser
from apps.kpi.service import executer
from apps.rapports.contenu import pptx_evenement, pptx_requete, xlsx_evenement
from apps.rapports.models import Rapport

from .test_evenements import base_horaire, evenement  # noqa: F401  (fixtures)
from .test_service import analyste, base_kpi, referentiel, requete  # noqa: F401

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def sans_modele(settings):
    """Charte dessinée par le code, sauf dans les tests qui fournissent un modèle."""
    settings.PPTX_MODELE = None


def _textes(prs):
    return [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text) for s in prs.slides]


def test_pptx_evenement(evenement, base_horaire, analyste):  # noqa: F811
    prs = Presentation(BytesIO(pptx_evenement(analyser(evenement, analyste, base_horaire))))
    textes = _textes(prs)
    assert textes[0].startswith("Foire")
    assert any("synthèse événement / référence" in t for t in textes)
    graphiques = [s for s in prs.slides if any(sh.has_chart for sh in s.shapes)]
    assert len(graphiques) == 4  # un par KPI analysé
    assert any("par rapport à la référence" in t for t in textes)
    assert any(t.startswith("Anomalies détectées") for t in textes)
    assert textes[-1].startswith("Annexe")


def test_pptx_requete(referentiel, base_kpi, analyste):  # noqa: F811
    prs = Presentation(BytesIO(pptx_requete(executer(requete(granularite_espace="commune"), analyste, base_kpi))))
    assert len([s for s in prs.slides if any(sh.has_chart for sh in s.shapes)]) == 2
    assert prs.slide_width > prs.slide_height


def test_pptx_avec_modele(evenement, base_horaire, analyste, tmp_path, settings):  # noqa: F811
    modele = Presentation()
    modele.slides.add_slide(modele.slide_layouts[0]).shapes.title.text = "DIAPO DU MODÈLE"
    modele.save(tmp_path / "modele.pptx")
    settings.PPTX_MODELE = tmp_path / "modele.pptx"
    prs = Presentation(BytesIO(pptx_evenement(analyser(evenement, analyste, base_horaire))))
    assert not any("DIAPO DU MODÈLE" in t for t in _textes(prs))
    assert prs.slides[0].slide_layout.name == "Title Slide"
    assert prs.slides[1].slide_layout.name == "Title Only"


def test_pptx_modele_helia(evenement, base_horaire, analyste, settings):  # noqa: F811
    from apps.rapports.powerpoint import charte_du_theme

    settings.PPTX_MODELE = settings.BASE_DIR / "config" / "modele_rapport.pptx"
    octets = pptx_evenement(analyser(evenement, analyste, base_horaire))
    prs = Presentation(BytesIO(octets))
    dispositions = [s.slide_layout.name for s in prs.slides]
    assert dispositions[0] == "Diapo 1" and dispositions[-1] == "Diapo FIN"
    assert set(dispositions[1:-1]) == {"Diapo Simple 3"}
    assert {l.name for l in prs.slide_layouts} == {"Diapo 1", "Diapo Simple 3", "Diapo FIN"}
    assert len(octets) < 4_000_000  # dispositions inutilisées (et leurs images) retirées
    # Titre écrit dans l'espace réservé du modèle, autres espaces réservés retirés.
    assert list(prs.slides[1].placeholders)[0].text_frame.text == "Périmètre analysé"
    assert list(prs.slides[2].placeholders)[0].text_frame.text.startswith("LTE — synthèse")
    assert len(prs.slides[1].placeholders) == 1
    assert prs.slides[0].shapes.title.text == "Foire"
    # Couleurs du thème : magenta pour l'événement, gris foncé pour la référence.
    charte = charte_du_theme(prs.slide_master.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme").blob)
    assert (charte.principale, charte.secondaire) == ("FF00E3", "5F6A72")
    graphe = next(sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart)
    assert str(graphe.series[0].format.line.color.rgb) == "FF00E3"


def test_xlsx_evenement(evenement, base_horaire, analyste):  # noqa: F811
    wb = load_workbook(BytesIO(xlsx_evenement(analyser(evenement, analyste, base_horaire, niveau="site"))))
    assert wb.sheetnames == ["Synthèse", "Anomalies", "Classement", "Courbes LTE"]
    assert wb["Anomalies"]["C1"].value == "Site"
    assert wb["Courbes LTE"].max_row == 1 + 5  # 5 heures d'événement


@pytest.fixture
def worker(monkeypatch, base_horaire, settings, tmp_path):  # noqa: F811
    """Exécute la tâche immédiatement (pas de worker django-q2 en test)."""
    from apps.rapports import taches

    settings.MEDIA_ROOT = tmp_path
    monkeypatch.setattr("apps.rapports.taches.moteur_kpi", lambda: base_horaire)
    monkeypatch.setattr("apps.rapports.lancement.async_task", lambda _nom, pk, **kw: taches.generer(pk))
    monkeypatch.setattr("apps.evenements.views.moteur_kpi", lambda: base_horaire)


@pytest.fixture
def client_analyste(client, analyste):  # noqa: F811
    client.force_login(analyste)
    return client


def test_ecran_evenement(client_analyste, evenement, worker):  # noqa: F811
    html = client_analyste.get(f"/evenements/{evenement.pk}/").content.decode()
    assert "Synthèse événement / référence" in html and "Anomalies détectées" in html
    assert 'id="courbes-LTE"' in html and "Saturation" in html
    html = client_analyste.get(f"/evenements/{evenement.pk}/?niveau=site").content.decode()
    assert "Sites les plus dégradés" in html
    assert "Foire" in client_analyste.get("/evenements/").content.decode()


def test_generation_et_telechargement(client_analyste, evenement, worker, analyste):  # noqa: F811
    r = client_analyste.post(f"/rapports/evenement/{evenement.pk}/", {"niveau": "secteur", "format": "pptx"})
    assert r.status_code == 302
    rapport = Rapport.objects.get()
    assert rapport.statut == "termine" and rapport.fichier.name.endswith(".pptx")
    assert JournalAudit.objects.filter(action="rapport_evenement_pptx").exists()
    assert "Télécharger" in client_analyste.get("/rapports/").content.decode()
    r = client_analyste.get(f"/rapports/{rapport.pk}/telecharger/")
    assert r.status_code == 200 and b"".join(r.streaming_content)[:2] == b"PK"

    autre = User.objects.create_superuser("autre")
    client_analyste.force_login(autre)
    assert client_analyste.get(f"/rapports/{rapport.pk}/telecharger/").status_code == 404


def test_rapport_en_erreur(client_analyste, referentiel, worker):  # noqa: F811
    from apps.evenements.models import Evenement

    e = Evenement.objects.create(nom="Sans date", cellules=["AAAe1"])
    client_analyste.post(f"/rapports/evenement/{e.pk}/", {"niveau": "secteur", "format": "xlsx"})
    rapport = Rapport.objects.get()
    assert rapport.statut == "erreur" and "créneau" in rapport.message


def test_rapport_requete(client_analyste, referentiel, base_kpi, worker, monkeypatch):  # noqa: F811
    monkeypatch.setattr("apps.rapports.taches.moteur_kpi", lambda: base_kpi)
    params = {"techno": "LTE", "perimetre_type": "commune", "perimetre_valeurs": "NOUMEA",
              "debut": "2026-09-01", "fin": "2026-09-02", "granularite_temps": "jour",
              "granularite_espace": "site", "fenetre_horaire": "journee", "kpis": ["lte_rrc_setup_sr"]}
    client_analyste.post("/rapports/requete/", params)
    rapport = Rapport.objects.get()
    assert rapport.statut == "termine" and rapport.nature == "requete"
    assert "NOUMEA" in rapport.titre
