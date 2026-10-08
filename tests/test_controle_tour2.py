"""Non-régression du contrôle indépendant (tour 2) : affichage, couverture des données,
bornes de période, export Excel, rapports, événements, démo, liens après IA."""

import io
import re
import time
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import pytest
from django.conf import settings as django_settings
from django.contrib.auth.models import User
from openpyxl import load_workbook
from pptx import Presentation

from apps.comptes.models import Perimetre
from apps.evenements.analyse import analyser
from apps.evenements.detection import ecart_significatif
from apps.evenements.models import Creneau, Evenement
from apps.kpi.catalogue import catalogue, catalogue_yaml, kpis_principaux
from apps.kpi.export_excel import LIGNES_MAX, construire
from apps.kpi.forms import RequeteForm
from apps.kpi.recherche import affichage
from apps.kpi.requete import RequeteKpi
from apps.kpi.service import RequeteRefusee, Resultat, ResultatTechno, couverture_donnees, executer
from apps.rapports.contenu import pptx_requete

from .test_recherche_ia import SUCCES, groq  # noqa: F401  (fixtures)
from .test_recherche_regles import J, ref_recherche  # noqa: F401
from .test_recherche_vues import admin, base_recherche  # noqa: F401

pytestmark = pytest.mark.django_db


def _html(r):
    assert r.status_code == 200, r.status_code
    return r.content.decode()


# ------------------------------------------------------------------ M3 / m6 / m7 / m13 / m4 : affichage

def test_formatage_des_cartes():
    texte, unite = affichage.formater(48_817_967, "appels")
    assert texte == "48 817 967" and unite == "appels"  # espace fine insécable
    assert affichage._taille(texte) == "longue" and affichage._taille("1 234 567 890") == "tres-longue"
    assert affichage.formater(4_947_560_878.1, "Mo") == ("4 948", "To")
    assert affichage.formater(2_500_000, "Mo", diviseur=1e6) == ("2,5", "To")  # même unité que la carte


def test_cartes_courbes_et_classement_dans_la_meme_unite(admin, base_recherche):  # noqa: F811
    html = _html(admin.get("/", {"q": "volume data 4G par site à Nouméa la semaine dernière"}))
    carte = re.search(r'<div class="valeur[^"]*"><span class="nombre-kpi">([^<]+)</span> <small class="unite">([^<]+)</small>',
                      html)
    assert carte and carte[2] in ("Go", "To", "Mo")
    donnees = re.search(r'id="graphiques-LTE"[^>]*>(.*?)</script>', html, re.S)[1]
    assert f'"unite": "{carte[2]}"' in donnees
    assert "Classement par site" not in html  # un volume ne se « dégrade » pas (m13)


def test_classement_qualite_seulement(admin, base_recherche):  # noqa: F811
    html = _html(admin.get("/", {"q": "drop 4G et volume data par site à Nouméa la semaine dernière"}))
    assert "Les plus dégradés — Taux de coupure E-RAB" in html
    assert "Les plus dégradés — Appels" not in html and "Durée moyenne" not in re.sub(
        r"<script.*?</script>", "", html, flags=re.S).split("Classement par site")[-1].split("Évolution")[0]


def test_tableau_detaille_tronque_annonce(admin, base_recherche, monkeypatch):  # noqa: F811
    monkeypatch.setattr(affichage, "LIGNES_MAX", 10)
    html = _html(admin.get("/", {"q": "drop 4G par cellule à Nouméa la semaine dernière"}))
    assert re.search(r"Tableau détaillé \(10 premières lignes sur \d+\)", html)


def test_causes_negligeables_seuil_coherent(admin, base_recherche):  # noqa: F811
    html = _html(admin.get("/", {"q": "causes de coupure 4G à Nouméa la semaine dernière"}))
    m = re.search(r"Négligeables \(&lt; ([\d,]+) % des coupures\)", html)
    assert m and m[1] == "0,1" and affichage.PART_NEGLIGEABLE == 0.1


def test_panneau_kpi_techno_demandee_et_coches_en_tete(admin, base_recherche):  # noqa: F811
    html = _html(admin.get("/", {"q": "drop 3G à Nouméa la semaine dernière"}))
    panneau = html.split('<div class="liste-kpi">')[1].split("</div>")[0]
    assert panneau.index("<legend>3G</legend>") < panneau.index("<legend>4G</legend>")
    assert panneau.index("Sélection actuelle") < panneau.index('value="wcdma_cs_drop" checked')
    assert panneau.index('value="wcdma_cs_drop" checked') < panneau.index('value="wcdma_hs_user_thp"')


# ------------------------------------------------------------------ M5 : couverture des données

def test_couverture_calculee_sur_les_donnees_lues():
    df = pd.DataFrame({"horodatage": pd.to_datetime(["2026-09-02 10:00", "2026-09-05 23:00"])})
    assert couverture_donnees(df, date(2026, 9, 2), date(2026, 9, 5)) is None
    assert couverture_donnees(df, date(2026, 9, 1), date(2026, 9, 10)) == (
        "données disponibles du 02/09/2026 au 05/09/2026 seulement (période demandée : du 01/09/2026 au 10/09/2026).")
    assert "aucune donnée" in couverture_donnees(df.iloc[0:0], date(2025, 1, 1), date(2025, 1, 2))


def test_avertissement_periode_hors_donnees(admin, base_recherche):  # noqa: F811
    html = _html(admin.get("/", {"q": "drop 4G Nouméa septembre 2025"}))
    assert "aucune donnée sur la période demandée (du 01/09/2025 au 30/09/2025)" in html
    html = _html(admin.get("/", {"q": "drop 4G Nouméa", "debut": "2026-09-20", "fin": "2026-10-07"}))
    assert "données disponibles du 24/09/2026 au 07/10/2026 seulement" in html  # démo : 14 jours


def test_bandeau_demo_periode_et_exemples(admin, base_recherche, monkeypatch):  # noqa: F811
    html = _html(admin.get("/"))
    assert "du 24/09/2026 au 07/10/2026" in html and "python manage.py charger_demo_kpi" in html
    assert "Causes de coupure 4G à Koné hier" in html  # démo à jour : « hier » disponible
    monkeypatch.setattr("apps.kpi.views.timezone.localdate", lambda: date(2026, 10, 20))
    html = _html(admin.get("/"))
    assert "Causes de coupure 4G à Koné le 07/10/2026" in html  # démo ancienne : dernier jour disponible


# ------------------------------------------------------------------ m1 : bornes de période

@pytest.mark.parametrize("params,message", [
    ({"debut": "0001-01-01", "fin": "9999-12-31"}, "Période hors limites"),
    ({"debut": "2025-01-01", "fin": "2026-10-07", "granularite_temps": "heure"}, "Période trop longue pour un calcul horaire"),
    ({"debut": "2005-01-01", "fin": "2026-10-07"}, "Période trop longue"),
])
def test_periodes_extremes_message_clair(admin, base_recherche, params, message):  # noqa: F811
    r = admin.get("/", {"q": "drop 4G Nouméa", **params})
    assert r.status_code == 200 and message in r.content.decode()
    r = admin.get("/", {"q": "drop 4G Nouméa", **params, "export": "xlsx"})
    assert r.status_code == 200 and message in r.content.decode()


def test_periode_longue_en_journalier_acceptee(base_recherche, ref_recherche):  # noqa: F811
    admin = User.objects.create_superuser("borne")
    req = RequeteKpi(techno=["LTE"], perimetre={"type": "global"},
                     periode={"debut": date(2017, 1, 1), "fin": date(2026, 10, 7)}, granularite_temps="mois",
                     kpis=["lte_erab_drop"])
    res = executer(req, admin, base_recherche)
    assert any("données disponibles du 24/09/2026" in a for a in res.avertissements)
    with pytest.raises(RequeteRefusee):
        executer(req.model_copy(update={"granularite_temps": "heure"}), admin, base_recherche)


# ------------------------------------------------------------------ M6 : export Excel

def _resultat(n_lignes: int, kpis=("lte_erab_drop", "lte_acces", "lte_payload_dl", "lte_dl_user_thp")):
    cat = catalogue()
    kpis = [cat[c] for c in kpis]
    periodes = pd.date_range("2026-09-01", periods=max(1, n_lignes // 50), freq="h")
    index = pd.MultiIndex.from_product([periodes, [f"C{i}" for i in range(50)]], names=["periode", "entite"])
    table = pd.DataFrame(np.random.default_rng(1).random((len(index), len(kpis))) * 3, index=index,
                         columns=[k.code for k in kpis])
    req = RequeteKpi(techno=["LTE"], perimetre={"type": "global"},
                     periode={"debut": date(2026, 9, 1), "fin": date(2026, 9, 30)}, granularite_temps="heure",
                     granularite_espace="cellule", kpis=[k.code for k in kpis])
    return Resultat(requete=req, par_techno=[ResultatTechno("LTE", kpis, table, 50, [],
                                                            table.groupby(level="entite").mean(),
                                                            {k.code: 1.0 for k in kpis})])


def test_export_excel_50000_lignes_en_quelques_secondes():
    t = time.monotonic()
    octets = construire(_resultat(50_000))
    assert time.monotonic() - t < 20
    wb = load_workbook(io.BytesIO(octets), read_only=True)
    assert wb["Données LTE"].max_row == 50_001
    parametres = {r[0]: r[1] for r in wb["Paramètres"].iter_rows(values_only=True)}
    assert parametres["Fenêtre horaire"] == "Journée complète" and parametres["Granularité temporelle"] == "par heure"


def test_export_excel_plafonne(monkeypatch):
    monkeypatch.setattr("apps.kpi.export_excel.LIGNES_MAX", 100)
    wb = load_workbook(io.BytesIO(construire(_resultat(500))))
    assert wb["Données LTE"].max_row == 101
    parametres = {r[0]: r[1] for r in wb["Paramètres"].iter_rows(values_only=True)}
    assert parametres["Données LTE"].startswith("onglet limité aux 100 premières lignes sur 500")
    assert LIGNES_MAX == 200_000


# ------------------------------------------------------------------ m12 : rapports

def test_pptx_causes_en_pareto_et_libelles_lisibles(base_recherche, ref_recherche):  # noqa: F811
    admin = User.objects.create_superuser("pptx")
    cat = catalogue()
    causes = [c for c, k in cat.items() if k.decomposition_de == "lte_erab_drop"]
    req = RequeteKpi(techno=["LTE"], perimetre={"type": "global"},
                     periode={"debut": date(2026, 9, 28), "fin": date(2026, 10, 4)}, kpis=["lte_erab_drop", *causes])
    prs = Presentation(io.BytesIO(pptx_requete(executer(req, admin, base_recherche))))
    textes = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert "Fenêtre horaire : Journée complète" in textes[0]
    graphiques = [s for s in prs.slides if any(sh.has_chart for sh in s.shapes)]
    assert len(graphiques) == 2  # courbe du drop + Pareto des causes (pas une diapo par cause)
    assert any("répartition par cause" in t for t in textes)


# ------------------------------------------------------------------ M4 : événements

def test_kpis_principaux_sans_causes_ni_composants():
    cat = catalogue_yaml()
    principaux = kpis_principaux(cat)
    assert "lte_erab_drop" in principaux and "lte_acces" in principaux
    assert not {"lte_drop_radio", "lte_rrc_setup_sr", "wcdma_rrc_cs_sr", "wcdma_drop_sho"} & set(principaux)


def test_ecart_absolu_minimal():
    cat = catalogue_yaml()
    cause = cat["lte_drop_transport"]
    assert cause.ecart_min == 0.1 and cat["lte_acces"].ecart_min == 0.5
    assert not ecart_significatif(cause, 0.02, 0.01, [0.01, 0.011, 0.009], 2, 20)  # +100 % mais 0,01 pt
    assert ecart_significatif(cat["lte_erab_drop"], 2.5, 1.7, [1.7, 1.69, 1.71], 2, 20)


def test_evenement_analyse_kpis_principaux(base_recherche, ref_recherche):  # noqa: F811
    tz = ZoneInfo(django_settings.TIME_ZONE)
    e = Evenement.objects.create(nom="Contrôle", sites=["NOU001"], semaines_reference=1)
    Creneau.objects.create(evenement=e, debut=datetime(2026, 10, 6, 7, tzinfo=tz), fin=datetime(2026, 10, 6, 11, tzinfo=tz))
    a = analyser(e, User.objects.create_superuser("evt"), base_recherche)
    codes = {k.code for r in a.par_techno for k in r.kpis}
    assert codes == set(kpis_principaux(catalogue()))
    assert not any(an.kpi and an.kpi.decomposition_de for an in a.anomalies)
    e.kpis = ["lte_drop_radio"]  # une cause reste sélectionnable explicitement
    e.save()
    assert [k.code for r in analyser(e, User.objects.get(username="evt"), base_recherche).par_techno
            for k in r.kpis] == ["lte_drop_radio"]


# ------------------------------------------------------------------ m10 / m11 / m16

def test_exemples_du_lecteur_restreint(client, base_recherche):  # noqa: F811
    lecteur = User.objects.create_user("lec")
    Perimetre.objects.create(nom="Païta", communes=["PAITA"]).utilisateurs.add(lecteur)
    client.force_login(lecteur)
    html = _html(client.get("/"))
    exemples = re.findall(r'<a class="exemple"[^>]*>([^<]+)</a>', html)
    assert exemples and all("Nouméa" not in e and "Koné" not in e and "Lifou" not in e for e in exemples)
    assert sum("Païta" in e for e in exemples) >= 4


def test_liens_apres_recherche_ia_sans_nouvel_appel(admin, base_recherche, groq, settings):  # noqa: F811
    groq.reponse = SUCCES
    html = _html(admin.get("/", {"q": "les coupures 4G le soir à Nouméa la semaine dernière", "ia": "1"}))
    assert len(groq.appels) == 1
    lien = re.search(r'href="(\?[^"]*export=xlsx[^"]*)"', html)[1]
    assert "ia=1" not in lien and "perimetre_valeurs=NOUMEA" in lien and "debut=2026-09-28" in lien
    caches = re.findall(r'<input type="hidden" name="(\w+)" value="([^"]*)">', html.split('class="compris"')[1])
    assert ("ia", "1") not in caches and ("fenetre_horaire", "18-22") in caches
    r = admin.get("/" + lien.replace("&amp;", "&"))
    assert r.status_code == 200 and len(groq.appels) == 1  # export : pas de nouvel appel au modèle


def test_formulaire_dates_par_defaut_heure_locale(monkeypatch):
    monkeypatch.setattr("apps.kpi.forms.timezone.localdate", lambda: date(2026, 10, 8))
    form = RequeteForm()
    assert form.fields["fin"].initial() == date(2026, 10, 7)
    assert form.fields["debut"].initial() == date(2026, 10, 1)


# ------------------------------------------------------------------ m15 : générateur

def test_generateur_causes_plafonnees_ligne_a_ligne(tmp_path):
    from apps.kpi.management.commands.charger_demo_kpi import TECHNOS, Generateur, lire_bases
    from apps.kpi.service import colonnes_utilisees

    colonnes = colonnes_utilisees(list(catalogue_yaml().values()))
    bases = lire_bases(django_settings.BASE_DIR, colonnes, cellules_max=40)
    rng = np.random.default_rng(5)
    for techno, (total, causes_rx) in {"LTE": ("ErabDropRate_p", "ErabDrop"), "WCDMA": ("RabDropCs_p", "RabDropCs")}.items():
        gen = Generateur(bases[techno], TECHNOS[techno], colonnes, rng)
        gen.proba = {c: np.minimum(p * 40, 1) for c, p in gen.proba.items()}  # taux extrêmes : plafond sollicité
        if gen.reste is not None:
            gen.reste = np.minimum(gen.reste * 40, 1)
        v = gen.jour(date(2026, 9, 30), [])
        causes = [c for c in gen.parts]
        somme = sum(v[c] for c in causes)
        assert (somme <= v[total] + 1e-9).all() and (v[total] <= 100 + 1e-9).all()
        if techno == "LTE":
            assert np.allclose(somme, v[total])
