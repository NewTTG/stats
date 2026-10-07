"""Export Excel : onglets, valeurs de synthèse (ratio de sommes) et seuils."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.kpi.catalogue import catalogue
from apps.kpi.export_excel import classeur, nom_fichier
from apps.kpi.service import executer

from .test_service import analyste, base_kpi, referentiel, requete  # noqa: F401  (fixtures)

pytestmark = pytest.mark.django_db


def test_statut_selon_sens_et_seuils():
    rrc, drop = catalogue()["lte_rrc_setup_sr"], catalogue()["lte_erab_drop"]
    assert rrc.statut(99.5) == "" and rrc.statut(97) == "alerte" and rrc.statut(90) == "critique"
    assert drop.statut(0.5) == "" and drop.statut(1.5) == "alerte" and drop.statut(3) == "critique"
    assert rrc.statut(None) == "" and catalogue()["lte_payload_dl"].statut(1) == ""


def test_synthese_ratio_de_sommes_sur_la_periode(referentiel, base_kpi, analyste):  # noqa: F811
    res = executer(requete(), analyste, base_kpi).par_techno[0]
    assert res.synthese["lte_rrc_setup_sr"] == pytest.approx(1059 * 4 / (1160 * 4) * 100)
    assert res.synthese["lte_payload_dl"] == pytest.approx(160.0)


def test_classeur(referentiel, base_kpi, analyste):  # noqa: F811
    req = requete(perimetre={"type": "commune", "valeurs": ["PAITA", "NOUMEA"]}, granularite_espace="site")
    resultat = executer(req, analyste, base_kpi)
    wb = load_workbook(BytesIO(classeur(resultat)))
    assert wb.sheetnames == ["Synthèse", "Données LTE", "Cellules sans données"]

    synthese = {r[1]: r for r in wb["Synthèse"].iter_rows(values_only=True) if r[0] == "LTE"}
    rrc = synthese["Taux de succès d'établissement RRC"]
    assert rrc[3] == pytest.approx(1009 * 4 / (1110 * 4) * 100)
    assert rrc[4] == "Critique"  # 90,9 % < 95 %

    donnees = list(wb["Données LTE"].iter_rows(values_only=True))
    assert donnees[0][:2] == ("Période", "Entité")
    assert len(donnees) == 1 + 2 * 2  # 2 jours × 2 sites
    assert donnees[1][0].date() == date(2026, 9, 1)
    assert list(wb["Cellules sans données"].iter_rows(values_only=True))[1] == ("LTE", "BBBe2")

    assert nom_fichier(resultat) == "kpi_lte_20260901_20260902.xlsx"
