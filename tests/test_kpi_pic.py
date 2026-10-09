"""KPI « pic » (utilisateurs connectés) : somme des cellules à chaque horodatage, puis maximum
sur le temps. Catalogue, moteur, requête, analyse d'événement, détection, démo, affichage."""

import io
import math
import sqlite3
from datetime import date, datetime

import numpy as np
import pandas as pd
import pytest
import sqlalchemy as sa
from django.conf import settings as django_settings
from django.contrib.auth.models import User
from django.utils import timezone
from openpyxl import load_workbook
from pptx import Presentation

from apps.evenements import detection
from apps.evenements.analyse import _valeurs, analyser
from apps.evenements.models import Creneau, Evenement
from apps.kpi.catalogue import DefinitionKpi, catalogue, catalogue_yaml, kpis_principaux
from apps.kpi.export_excel import construire
from apps.kpi.management.commands.charger_demo_kpi import PROFIL_SEMAINE, TECHNOS, _nature, lire_bases
from apps.kpi.moteur import agreger
from apps.kpi.recherche import affichage
from apps.kpi.requete import RequeteKpi
from apps.kpi.service import colonnes_utilisees, executer
from apps.rapports.contenu import pptx_evenement, pptx_requete, xlsx_evenement

from .test_demo_kpi import demo, extraits  # noqa: F401  (fixtures)
from .test_moteur import DROP, VOLUME, _catalogue
from .test_recherche_regles import J, ref_recherche  # noqa: F401
from .test_recherche_vues import admin, base_recherche  # noqa: F401
from .test_service import analyste, referentiel  # noqa: F401

pytestmark = pytest.mark.django_db

CODE = "lte_rrc_conn_max"
PIC = DefinitionKpi(code="uc", libelle="Utilisateurs", techno="LTE", unite="utilisateurs", categorie="trafic",
                    numerateur="rrc", agregation="pic", sens="haut_est_mieux")


def _df():
    """2 cellules × 3 heures × 2 jours ; le 01/09, A culmine à 10 h et B à 11 h."""
    valeurs = {("A", 1): [10, 2, 3], ("B", 1): [1, 8, 4], ("A", 2): [20, 1, 1], ("B", 2): [1, 1, 1]}
    lignes = [{"cellule": c, "horodatage": datetime(2026, 9, jour, 10 + i), "rrc": float(v),
               "drops": 1, "releases": 100, "volume": 1.0}
              for (c, jour), serie in valeurs.items() for i, v in enumerate(serie)]
    df = pd.DataFrame(lignes)
    return df.assign(entite=df["cellule"], periode=df["horodatage"].dt.normalize())


# ------------------------------------------------------------------ moteur

def test_pic_somme_par_instant_puis_maximum():
    res = agreger(_df(), [PIC], par=[]).iloc[0]
    # 01/09 : Σ par heure 11, 10, 7 ; 02/09 : 21, 2, 2 -> 21 (≠ Σ des pics par cellule : 20 + 8).
    assert res["uc"] == 21 and res["uc__num"] == 21 and res["uc__den"] == 1
    premier_jour = agreger(_df()[lambda d: d["horodatage"].dt.day == 1], [PIC], par=[])["uc"].iloc[0]
    assert premier_jour == 11 and premier_jour != 10 + 8


def test_pic_par_entite_et_par_periode():
    df = _df()
    par_cellule = agreger(df, [PIC], par=["entite"])["uc"]
    assert par_cellule.to_dict() == {"A": 20, "B": 8}
    par_jour = agreger(df, [PIC], par=["periode"])["uc"]
    assert list(par_jour) == [11, 21]
    detail = agreger(df, [PIC], par=["periode", "entite"])
    assert detail["uc"].to_dict() == {(pd.Timestamp("2026-09-01"), "A"): 10, (pd.Timestamp("2026-09-01"), "B"): 8,
                                      (pd.Timestamp("2026-09-02"), "A"): 20, (pd.Timestamp("2026-09-02"), "B"): 1}
    assert (detail["uc__den"] == 1).all()
    # Par horodatage : un seul instant par groupe, le pic est la somme des cellules.
    assert list(agreger(df, [PIC], par=["horodatage"])["uc"]) == [11, 10, 7, 21, 2, 2]


def test_pic_sans_horodatage_somme_simple():
    df = _df().drop(columns="horodatage")
    assert agreger(df, [PIC], par=[])["uc"].iloc[0] == df["rrc"].sum()


def test_pic_absence_de_donnees_distincte_de_zero():
    df = _df()
    df.loc[(df["cellule"] == "B") & (df["horodatage"] == datetime(2026, 9, 1, 10)), "rrc"] = np.nan
    df.loc[df["cellule"] == "A", "rrc"] = np.where(df.loc[df["cellule"] == "A", "periode"].dt.day == 2, np.nan, 0)
    res = agreger(df, [PIC], par=["periode", "entite"])["uc"]
    assert math.isnan(res[(pd.Timestamp("2026-09-02"), "A")])  # pas de données
    assert res[(pd.Timestamp("2026-09-01"), "A")] == 0  # valeur nulle réelle
    assert res[(pd.Timestamp("2026-09-01"), "B")] == 8  # heure manquante ignorée
    assert agreger(df, [PIC], par=["periode"])["uc"].tolist() == [8, 1]  # 10 h : A = 0, B absente
    assert math.isnan(agreger(df.assign(rrc=np.nan), [PIC], par=[])["uc"].iloc[0])


def test_pic_avec_d_autres_kpi():
    df = _df()
    res = agreger(df, [DROP, PIC, VOLUME], par=["entite"])
    assert res.loc["A", "drop"] == pytest.approx(1.0) and res.loc["A", "vol"] == 6
    assert res.loc["A", "uc"] == 20 and res.loc["B", "uc"] == 8


# ------------------------------------------------------------------ catalogue

def test_catalogue_utilisateurs_connectes():
    k = catalogue_yaml()[CODE]
    assert k.pic and k.additif and k.agregation == "pic" and k.denominateur is None
    assert k.techno == "LTE" and k.categorie == "trafic" and k.unite == "utilisateurs" and k.qualite == "approx"
    assert k.formule == "max_t Σ RrcConnMax_max_nb" and "majorant" in k.note
    assert colonnes_utilisees([k]) == ["RrcConnMax_max_nb"]
    assert CODE in kpis_principaux(catalogue_yaml())
    volume = catalogue_yaml()["lte_payload_dl"]
    assert not volume.pic and volume.agregation == "somme" and volume.formule == "Σ PayloadDl_mB"


@pytest.mark.parametrize("modif", [
    {"denominateur": "x"},
    {"produit_de": ["a", "b"], "numerateur": None},
    {"decomposition_de": "parent"},
])
def test_pic_numerateur_seul(modif):
    with pytest.raises(ValueError, match="pic"):
        DefinitionKpi(**{**PIC.model_dump(exclude={"composants"}), **modif})


def test_pic_ni_composant_ni_agregation_inconnue(tmp_path):
    rrc = {"code": "rrc", "libelle": "RRC", "techno": "LTE", "unite": "%", "categorie": "accessibilite",
           "numerateur": "s", "denominateur": "a", "facteur": 100, "sens": "haut_est_mieux"}
    pic = {**PIC.model_dump(exclude={"composants"}), "seuils": {}}
    composite = {**rrc, "code": "acces", "numerateur": None, "denominateur": None, "produit_de": ["rrc", "uc"]}
    with pytest.raises(ValueError, match="additif"):
        _catalogue(tmp_path, [rrc, pic, composite])
    with pytest.raises(ValueError, match="agregation"):
        DefinitionKpi(**{**pic, "agregation": "moyenne"})


# ------------------------------------------------------------------ analyse d'événement, détection

def test_reference_evenement_moyenne_des_pics_hebdomadaires():
    instants = [datetime(2026, 9, 29, 18), datetime(2026, 9, 29, 19)]
    valeurs = {0: ([1, 1], [1, 2]), 1: ([10, 0], [0, 4]), 2: ([4, 4], [4, 4])}  # semaine : (A, B) par heure
    df = pd.DataFrame([{"tout": "Global", "cellule": c, "semaine": s, "instant": t,
                        "horodatage": t - pd.Timedelta(weeks=s), "rrc": serie[i]}
                       for s, paire in valeurs.items() for c, serie in zip("AB", paire)
                       for i, t in enumerate(instants)])
    valeur, reference, hebdo = _valeurs(df, [PIC], ["tout"])
    assert valeur.loc["Global", "uc"] == 3  # Σ par heure 2, 3
    assert list(hebdo["uc"]) == [10, 8]  # pics hebdomadaires
    assert reference.loc["Global", "uc"] == 9  # moyenne des pics, pas leur maximum (10)
    # Courbes : Σ des cellules à chaque instant, référence = moyenne des semaines.
    valeur, reference, hebdo = _valeurs(df, [PIC], ["instant"])
    assert list(valeur["uc"]) == [2, 3] and list(reference["uc"]) == [9, 6]
    assert hebdo["uc"].unstack("semaine").values.tolist() == [[10, 8], [4, 8]]


def test_detection_pic_comme_un_volume():
    cat = catalogue_yaml()
    for code in (CODE, "lte_payload_dl"):
        k = cat[code]
        assert not detection.taux_de_succes(k)
        assert detection.ecart_significatif(k, 3, 9, [10, 8], sigma=2, pct=20)  # baisse : dégradation
        assert not detection.ecart_significatif(k, 30, 9, [10, 8], sigma=2, pct=20)  # hausse : non
    assert affichage._volume(cat[CODE]) and not affichage._qualite(cat[CODE])


@pytest.fixture
def base_pic():
    """AAAe1 / AAAe2 (SITE_AAA), BBBe1 (SITE_BBB) : 18 h et 19 h, les 15, 22 et 29/09 ; 10 h et 11 h, les 01-02/09."""
    horaire = {
        (29, "AAAe1"): [1, 1], (29, "AAAe2"): [1, 2], (29, "BBBe1"): [5, 5],
        (22, "AAAe1"): [10, 0], (22, "AAAe2"): [0, 4], (22, "BBBe1"): [5, 5],
        (15, "AAAe1"): [4, 4], (15, "AAAe2"): [4, 4], (15, "BBBe1"): [5, 5],
    }
    lignes = [{"EutranCell_Id": c, "DateHour": datetime(2026, 9, j, 18 + i), "RrcConnMax_max_nb": v}
              for (j, c), serie in horaire.items() for i, v in enumerate(serie)]
    requete = {(1, "AAAe1"): [10, 2], (1, "AAAe2"): [1, 8], (1, "BBBe1"): [3, 3],
               (2, "AAAe1"): [4, 6], (2, "AAAe2"): [5, 5], (2, "BBBe1"): [0, 9]}
    lignes += [{"EutranCell_Id": c, "DateHour": datetime(2026, 9, j, 10 + i), "RrcConnMax_max_nb": v}
               for (j, c), serie in requete.items() for i, v in enumerate(serie)]
    h = pd.DataFrame(lignes)
    jour = (h.assign(DateDay=h["DateHour"].dt.date).groupby(["EutranCell_Id", "DateDay"], as_index=False)
            ["RrcConnMax_max_nb"].max())
    engine = sa.create_engine("sqlite://", poolclass=sa.pool.StaticPool)
    h.to_sql("lte_cell_hour", engine, index=False)
    jour.to_sql("lte_cell_day", engine, index=False)
    return engine


def _requete(**modifs):
    return RequeteKpi(**{"techno": ["LTE"], "perimetre": {"type": "global"},
                         "periode": {"debut": date(2026, 9, 1), "fin": date(2026, 9, 2)},
                         "granularite_temps": "jour", "granularite_espace": "site", "kpis": [CODE], **modifs})


def test_requete_pic_horaire_et_journaliere(referentiel, base_pic, analyste):  # noqa: F811
    # Lecture horaire (fenêtre 0-24) : Σ des cellules à chaque heure, maximum sur la journée.
    res = executer(_requete(fenetre_horaire="0-24"), analyste, base_pic).par_techno[0]
    assert res.table[CODE].to_dict() == {(pd.Timestamp("2026-09-01"), "SITE_AAA"): 11,
                                         (pd.Timestamp("2026-09-01"), "SITE_BBB"): 3,
                                         (pd.Timestamp("2026-09-02"), "SITE_AAA"): 11,
                                         (pd.Timestamp("2026-09-02"), "SITE_BBB"): 9}
    assert res.synthese[CODE].to_dict() == {"SITE_AAA": 11, "SITE_BBB": 9}
    assert res.synthese_globale[CODE] == 20  # 02/09 11 h : 6 + 5 + 9
    # Lecture journalière : Σ des pics journaliers des cellules, majorant de la valeur horaire.
    res = executer(_requete(), analyste, base_pic).par_techno[0]
    assert res.synthese[CODE].to_dict() == {"SITE_AAA": 18, "SITE_BBB": 9}
    assert res.synthese_globale[CODE] == 21  # 01/09 : 10 + 8 + 3
    # Par heure : la somme des cellules.
    res = executer(_requete(granularite_temps="heure", granularite_espace="global"), analyste, base_pic).par_techno[0]
    assert list(res.table[CODE]) == [14, 13, 9, 20]


def test_analyse_evenement_pic(referentiel, base_pic, analyste):  # noqa: F811
    e = Evenement.objects.create(nom="Concert", sites=["AAA100"], kpis=[CODE], semaines_reference=2)
    Creneau.objects.create(evenement=e, debut=timezone.make_aware(datetime(2026, 9, 29, 18)),
                           fin=timezone.make_aware(datetime(2026, 9, 29, 20)))
    a = analyser(e, analyste, base_pic, niveau="site")
    lte = a.par_techno[0]
    ligne = lte.synthese[0]
    assert ligne.kpi.code == CODE and ligne.valeur == 3 and ligne.reference == 9 and ligne.significatif
    courbe = lte.courbes[0]
    assert courbe.evenement == [2, 3] and courbe.reference == [9, 6]
    assert (courbe.ref_min, courbe.ref_max) == ([8, 4], [10, 8])
    ecart = next(x for x in a.anomalies if x.regle == "ecart")
    assert ecart.entite == "SITE_AAA" and ecart.kpi.code == CODE and ecart.gravite == "critique"
    # Rapports d'événement.
    assert Presentation(io.BytesIO(pptx_evenement(a))).slides
    wb = load_workbook(io.BytesIO(xlsx_evenement(a)))
    assert [c.value for c in wb["Synthèse"][5]][1:5] == ["Utilisateurs connectés (pic RRC)", "utilisateurs", 3, 9]


# ------------------------------------------------------------------ base de démonstration

def test_demo_pic_journalier_maximum_des_heures(demo):  # noqa: F811
    chemin, _ = demo
    with sqlite3.connect(chemin) as conn:
        h = pd.read_sql_query("SELECT EutranCell_Id, DateHour, RrcConnMax_max_nb FROM lte_cell_hour", conn)
        j = pd.read_sql_query("SELECT EutranCell_Id, DateDay, RrcConnMax_max_nb FROM lte_cell_day", conn)
        types = {nom: t for _, nom, t, *_ in conn.execute('PRAGMA table_info("lte_cell_hour")')}
    assert types["RrcConnMax_max_nb"] == "INTEGER"
    assert h["RrcConnMax_max_nb"].notna().all() and (h["RrcConnMax_max_nb"] >= 0).all()
    assert (h["RrcConnMax_max_nb"] == h["RrcConnMax_max_nb"].round()).all()
    h["DateDay"] = h["DateHour"].str[:10]
    maxima = h.groupby(["EutranCell_Id", "DateDay"])["RrcConnMax_max_nb"].max()
    jour = j.set_index(["EutranCell_Id", "DateDay"])["RrcConnMax_max_nb"].reindex(maxima.index)
    pd.testing.assert_series_equal(maxima, jour, check_dtype=False)
    par_heure = h.groupby(h["DateHour"].str[11:13])["RrcConnMax_max_nb"].sum()
    assert par_heure["19"] > 2 * par_heure["03"]  # profil horaire


def test_demo_pic_base_depuis_l_extrait_horaire(extraits):  # noqa: F811
    assert _nature("RrcConnMax_max_nb", pd.Series([10.0]), TECHNOS["LTE"]) == "pic"
    base = lire_bases(extraits, colonnes_utilisees(list(catalogue_yaml().values())))["LTE"]
    horaire = pd.read_csv(extraits / "lte_cell_hour.csv", sep=";").drop_duplicates("EutranCell_Id")
    horaire = horaire.set_index("EutranCell_Id")
    cellules = base.index[base["_origine"] == "heure"]
    assert len(cellules)
    # Pic de minuit ramené au pic de la journée par la forme du profil ; compteurs par le poids de minuit.
    forme = PROFIL_SEMAINE[0] / PROFIL_SEMAINE.max()
    assert np.allclose(base.loc[cellules, "RrcConnMax_max_nb"], horaire.loc[cellules, "RrcConnMax_max_nb"] / forme)
    assert np.allclose(base.loc[cellules, "PayloadDl_mB"], horaire.loc[cellules, "PayloadDl_mB"] / PROFIL_SEMAINE[0])


# ------------------------------------------------------------------ affichage et exports

def test_affichage_et_exports(admin, base_recherche):  # noqa: F811
    params = {"techno": "LTE", "perimetre_type": "global", "debut": "2026-10-01", "fin": "2026-10-07",
              "granularite_temps": "jour", "granularite_espace": "site", "fenetre_horaire": "journee",
              "kpis": [CODE, "lte_payload_dl"]}
    r = admin.get("/", params)
    html = r.content.decode()
    assert r.status_code == 200 and "Les plus chargés — Utilisateurs connectés (pic RRC)" in html

    req = RequeteKpi(techno=["LTE"], perimetre={"type": "global"},
                     periode={"debut": date(2026, 10, 1), "fin": date(2026, 10, 7)},
                     granularite_espace="site", kpis=[CODE, "lte_payload_dl"])
    resultat = executer(req, User.objects.get(username="admin"), base_recherche)
    res = resultat.par_techno[0]
    with sqlite3.connect(django_settings.KPI_DEMO_SQLITE) as conn:
        par_jour = pd.read_sql_query("SELECT DateDay, SUM(RrcConnMax_max_nb) AS s FROM lte_cell_day "
                                     "WHERE DateDay BETWEEN '2026-10-01' AND '2026-10-07' GROUP BY DateDay", conn)
    assert res.synthese_globale[CODE] == par_jour["s"].max() > 0

    carte = next(c for c in affichage.cartes(res) if c["kpi"].code == CODE)
    assert carte["unite"] == "utilisateurs" and carte["approx"] and "," not in carte["texte"]
    graphe = next(g for g in affichage.graphiques(res, "jour")["kpis"] if g["code"] == CODE)
    assert graphe["zero"]
    classement = next(c for c in affichage.classement(res) if c["kpi"].code == CODE)
    valeurs = [ligne["valeur"] for ligne in classement["lignes"]]
    assert classement["titre"] == "Les plus chargés" and valeurs == sorted(valeurs, reverse=True)

    wb = load_workbook(io.BytesIO(construire(resultat)))
    assert "Utilisateurs connectés (pic RRC) (utilisateurs) ≈" in [c.value for c in wb["Synthèse LTE"][1]]
    definitions = {ligne[0]: ligne for ligne in wb["Définitions"].iter_rows(values_only=True)}
    assert definitions[CODE][3] == "max_t Σ RrcConnMax_max_nb"
    prs = Presentation(io.BytesIO(pptx_requete(resultat)))
    textes = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame) for s in prs.slides]
    assert any("Utilisateurs connectés (pic RRC)" in t for t in textes)
    assert catalogue()[CODE].statut(res.synthese_globale[CODE]) == ""  # pas de seuil
