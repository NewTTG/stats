import pytest
from django.contrib.auth.models import User

from apps.comptes.models import JournalAudit

from .test_service import base_kpi, referentiel  # noqa: F401  (fixtures)

pytestmark = pytest.mark.django_db

PARAMS = {
    "techno": "LTE", "perimetre_type": "commune", "perimetre_valeurs": "NOUMEA",
    "debut": "2026-09-01", "fin": "2026-09-02", "granularite_temps": "jour",
    "granularite_espace": "site", "fenetre_horaire": "journee", "kpis": ["lte_rrc_setup_sr"],
}


@pytest.fixture
def client_admin(client):
    client.force_login(User.objects.create_superuser("admin"))
    return client


def test_connexion_requise_sans_acces_anonyme(client, settings):
    settings.ACCES_ANONYME, settings.LOGIN_URL = False, "login"
    r = client.get("/")
    assert r.status_code == 302 and r.url.startswith("/connexion/")


def test_formulaire_vierge(client_admin):
    r = client_admin.get("/")
    html = r.content.decode()
    assert r.status_code == 200 and 'name="q"' in html  # barre de recherche
    assert "Recherche avancée" in html and 'name="perimetre_type"' in html  # formulaire structuré conservé


def test_resultat_et_audit(client_admin, referentiel, base_kpi, monkeypatch):  # noqa: F811
    monkeypatch.setattr("apps.kpi.views.moteur_kpi", lambda: base_kpi)
    r = client_admin.get("/", PARAMS)
    html = r.content.decode()
    assert "SITE_AAA" in html and "01/09/2026" in html
    assert JournalAudit.objects.filter(action="requete_kpi").count() == 1


def test_resultat_synthese_et_graphiques(client_admin, referentiel, base_kpi, monkeypatch):  # noqa: F811
    monkeypatch.setattr("apps.kpi.views.moteur_kpi", lambda: base_kpi)
    html = client_admin.get("/", PARAMS).content.decode()
    assert "Synthèse sur la période" in html
    assert 'id="graphiques-LTE"' in html and "echarts" in html
    assert "export=xlsx" in html
    assert "Rapport PowerPoint" in html and 'name="perimetre_valeurs" value="NOUMEA"' in html


def test_base_kpi_non_configuree(client_admin, referentiel, settings, tmp_path):  # noqa: F811
    from apps.kpi.source import moteur_kpi
    moteur_kpi.cache_clear()
    settings.KPI_DB = {**settings.KPI_DB, "host": ""}
    settings.KPI_DEMO_SQLITE = tmp_path / "absente.sqlite3"  # ni base KPI ni base de démonstration
    r = client_admin.get("/", PARAMS)
    assert "Base KPI non configurée" in r.content.decode()
    moteur_kpi.cache_clear()


def test_export_excel(client_admin, referentiel, base_kpi, monkeypatch):  # noqa: F811
    import io

    from openpyxl import load_workbook

    monkeypatch.setattr("apps.kpi.views.moteur_kpi", lambda: base_kpi)
    html = client_admin.get("/", PARAMS).content.decode()
    assert "export=xlsx" in html

    r = client_admin.get("/", {**PARAMS, "export": "xlsx"})
    assert r.status_code == 200
    assert r["Content-Disposition"] == 'attachment; filename="kpi_20260901_20260902.xlsx"'
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Paramètres", "Synthèse LTE", "Données LTE", "Définitions"]
    donnees = list(wb["Données LTE"].iter_rows(values_only=True))
    assert donnees[0][:2] == ("Période", "Entité")
    assert any(l[1] == "SITE_AAA" for l in donnees[1:])
    assert wb["Définitions"]["A2"].value == "lte_rrc_setup_sr"
    assert JournalAudit.objects.filter(action="export_excel").count() == 1


def test_export_excel_texte_pas_formule():
    from openpyxl import Workbook

    from apps.kpi.export_excel import _ajouter

    ws = Workbook().active
    _ajouter(ws, ["=HYPERLINK(\"x\")", 1.5])
    assert ws["A1"].data_type == "s" and ws["B1"].data_type == "n"
