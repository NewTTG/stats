"""Recherche en langage libre de bout en bout : vues, questions, exports, droits, bandeau démo.

Base KPI : base de démonstration générée sur les cellules du référentiel synthétique.
"""

import html as html_lib
import io
import re
from datetime import date

import pandas as pd
import pytest
from django.conf import settings as django_settings
from django.contrib.auth.models import User
from openpyxl import load_workbook

from apps.comptes.models import JournalAudit, Perimetre
from apps.kpi.management.commands.charger_demo_kpi import TECHNOS, generer
from apps.kpi.source import moteur_kpi
from apps.rapports.models import Rapport
from apps.referentiel.models import Cellule

from .test_recherche_regles import J, ref_recherche  # noqa: F401  (fixtures)

pytestmark = pytest.mark.django_db


@pytest.fixture
def base_recherche(ref_recherche, tmp_path, settings, monkeypatch):  # noqa: F811
    """Base de démonstration sur les cellules du référentiel synthétique, date figée (J)."""
    for t in TECHNOS.values():
        cellules = list(Cellule.objects.filter(techno=t.nom).values_list("nom", flat=True))
        for nom in (t.csv_heure, t.csv_jour):
            df = pd.read_csv(django_settings.BASE_DIR / nom, sep=";", dtype=str, nrows=200)
            df = df[pd.to_numeric(df[t.volume]) > 0].head(len(cellules)).copy()  # cellules actives
            df[t.cle] = cellules
            df.to_csv(tmp_path / nom, sep=";", index=False)
    sortie = tmp_path / "demo.sqlite3"
    generer(sortie, jours=14, graine=3, fin=date(2026, 10, 7), dossier=tmp_path, journal=lambda m: None)
    settings.KPI_DB = {**settings.KPI_DB, "host": ""}
    settings.KPI_DEMO_SQLITE = sortie
    settings.GROQ_API_KEY = ""
    monkeypatch.setattr("apps.kpi.views.timezone.localdate", lambda: J)
    moteur_kpi.cache_clear()
    yield moteur_kpi()
    moteur_kpi.cache_clear()


@pytest.fixture
def admin(client):
    client.force_login(User.objects.create_superuser("admin"))
    return client


def _html(reponse):
    assert reponse.status_code == 200
    return reponse.content.decode()


def _liens(html, texte):
    """Liens (href) des options dont le libellé commence par ``texte``."""
    return [html_lib.unescape(m[1]) for m in re.finditer(r'<a class="option" href="([^"]+)">' + re.escape(texte), html)]


def test_recherche_complete_affiche_le_resultat(admin, base_recherche):
    html = _html(admin.get("/", {"q": "drop 3G à Nouméa la semaine dernière"}))
    assert "Synthèse sur la période" in html and "Taux de coupure appels voix (3G)" in html
    assert "Nouméa" in html and "Semaine dernière · 28/09 → 04/10/2026" in html
    assert 'id="graphiques-WCDMA"' in html and "echarts" in html
    assert "export=xlsx" in html and "Rapport PowerPoint" in html
    assert 'name="perimetre_valeurs" value="NOUMEA"' in html and 'name="debut" value="2026-09-28"' in html
    assert "Sur quelle période" not in html
    trace = JournalAudit.objects.get(action="recherche_kpi").requete
    assert trace["q"] == "drop 3G à Nouméa la semaine dernière" and trace["ia"] is False
    assert trace["source"] == "regles" and trace["requete"]["kpis"] == ["wcdma_cs_drop", "wcdma_ps_drop"]


def test_recherche_incomplete_pose_une_question_puis_option(admin, base_recherche):
    html = _html(admin.get("/", {"q": "taux d'accès 4G Païta"}))
    assert "Sur quelle période ?" in html and "Synthèse sur la période" not in html
    assert not JournalAudit.objects.exists()  # rien n'est exécuté
    lien = _liens(html, "7 derniers jours")[0]
    assert "periode=7j" in lien and "q=taux" in lien
    html = _html(admin.get("/" + lien))
    assert "Synthèse sur la période" in html and "Taux d&#x27;accès 4G" in html
    assert "7 derniers jours · 01/10 → 07/10/2026" in html
    assert JournalAudit.objects.get().requete["parametres"] == {"periode": "7j"}


def test_question_quoi_et_dates_libres(admin, base_recherche):
    html = _html(admin.get("/", {"q": "Païta hier"}))
    assert "Que voulez-vous voir ?" in html
    lien = _liens(html, "Débit")[0]
    html = _html(admin.get("/" + lien))
    assert "Débit DL utilisateur moyen" in html and "Débit HSDPA utilisateur moyen" in html
    # Dates libres : réponse par les champs debut / fin.
    html = _html(admin.get("/", {"q": "débit 4G Païta", "debut": "2026-10-01", "fin": "2026-10-02"}))
    assert "Synthèse sur la période" in html and "01/10 → 02/10/2026" in html


def test_ambiguite_de_trigramme(admin, base_recherche):
    html = _html(admin.get("/", {"q": "débit 4G CHT hier"}))
    assert "Plusieurs lieux correspondent à « CHT »" in html
    assert len(_liens(html, "CHT")) == 3 and len(_liens(html, "Tous les sites CHT")) == 1
    lien = _liens(html, "CHT_1 (CHT802)")[0]
    assert "perimetre_type=site" in lien and "perimetre_valeurs=CHT802" in lien
    html = _html(admin.get("/" + lien))
    assert "Synthèse sur la période" in html and "CHT802" in html


def test_parametres_explicites_priment_sur_le_texte(admin, base_recherche):
    html = _html(admin.get("/", {"q": "drop 3G à Nouméa la semaine dernière", "techno": "LTE",
                                 "granularite_espace": "site"}))
    assert 'id="graphiques-LTE"' in html and 'id="graphiques-WCDMA"' not in html
    assert "Classement par site" in html


def test_causes_avec_non_ventile(admin, base_recherche):
    html = _html(admin.get("/", {"q": "causes des coupures voix 3G à Nouméa la semaine dernière"}))
    assert "Répartition par cause" in html and "Non ventilé" in html
    html = _html(admin.get("/", {"q": "causes de coupure 4G à Nouméa la semaine dernière"}))
    assert "Répartition par cause" in html and "Non ventilé" not in html  # causes LTE = total


def test_puces_modifiables_sans_javascript(admin, base_recherche):
    html = _html(admin.get("/", {"q": "drop 3G à Nouméa la semaine dernière", "ia": "1"}))
    assert html.count('<details class="puce') == 7
    # Chaque puce est un formulaire GET qui conserve la demande et remplace son seul champ.
    assert '<input type="hidden" name="q" value="drop 3G à Nouméa la semaine dernière">' in html
    assert '<button class="bouton secondaire petit" type="submit" name="periode" value="mois_dernier">' in html


def test_export_excel_depuis_une_recherche(admin, base_recherche):
    r = admin.get("/", {"q": "drop 3G à Nouméa la semaine dernière", "export": "xlsx"})
    assert r.status_code == 200 and r["Content-Disposition"] == 'attachment; filename="kpi_20260928_20261004.xlsx"'
    assert "Données WCDMA" in load_workbook(io.BytesIO(r.content)).sheetnames
    assert JournalAudit.objects.get(action="export_excel").requete["q"].startswith("drop 3G")


def test_rapport_powerpoint_depuis_une_recherche(admin, base_recherche, monkeypatch):
    monkeypatch.setattr("apps.rapports.views.async_task", lambda *a, **k: None)
    html = _html(admin.get("/", {"q": "drop 3G à Nouméa en soirée la semaine dernière"}))
    champs = re.search(r'action="/rapports/requete/">(.*?)</form>', html, re.S)[1]
    donnees = {}
    for nom, valeur in re.findall(r'<input type="hidden" name="(\w+)" value="([^"]*)">', champs):
        donnees.setdefault(nom, []).append(html_lib.unescape(valeur))
    assert admin.post("/rapports/requete/", donnees).status_code == 302
    requete = Rapport.objects.get().parametres["requete"]
    assert requete["fenetre_horaire"] == "18-22" and requete["perimetre"]["valeurs"] == ["NOUMEA"]


def test_lecteur_restreint_aucune_fuite(client, base_recherche):
    lecteur = User.objects.create_user("lec")
    Perimetre.objects.create(nom="Païta", communes=["PAITA"]).utilisateurs.add(lecteur)
    client.force_login(lecteur)
    html = _html(client.get("/", {"q": "drop 4G à Nouméa hier"}))
    assert "hors de votre périmètre" in html and "Synthèse sur la période" in html
    for interdit in ("NOU001", "AGENCE_TELECOM", "CHT801", "CHT_1", "NOU001L1", "Carnaval"):
        assert interdit not in html, interdit
    assert "PAI001" in html or "PAITA" in html
    r = client.get("/suggestions/", {"q": "n"})
    assert all(s["valeur"] in ("PAITA", "PAI001") for s in r.json()["lieux"])
    assert client.get("/suggestions/", {"q": "cht"}).json()["lieux"] == []


def test_bandeau_demo(admin, base_recherche, settings, tmp_path):
    assert "Données de démonstration (synthétiques)" in _html(admin.get("/", {"q": "drop 4G Païta hier"}))
    assert "Incidents simulés" in _html(admin.get("/", {"q": "drop 4G Païta hier"}))
    settings.KPI_DEMO_SQLITE = tmp_path / "absente.sqlite3"
    moteur_kpi.cache_clear()
    html = _html(admin.get("/", {"q": "drop 4G Païta hier"}))
    assert "Données de démonstration" not in html and "charger_demo_kpi" in html


def test_accueil_exemples_et_dernieres_recherches(admin, base_recherche):
    html = _html(admin.get("/"))
    assert "Drop 3G à Nouméa la semaine dernière" in html and "Vos recherches apparaîtront ici" in html
    admin.get("/", {"q": "débit 4G Païta hier"})
    html = _html(admin.get("/"))
    assert "Mes dernières recherches" in html and "?q=d%C3%A9bit+4G+Pa%C3%AFta+hier" in html
    assert 'name="ia"' not in html  # pas de clé : pas de case « Recherche IA »


def test_recherche_vide_et_texte_incomprehensible(admin, base_recherche):
    assert "Que voulez-vous savoir" in _html(admin.get("/", {"q": ""}))
    html = _html(admin.get("/", {"q": "xyzzy plop"}))
    assert "Mots non compris" in html and "xyzzy" in html
    assert "Que voulez-vous voir ?" in html and "Sur quelle période ?" in html
